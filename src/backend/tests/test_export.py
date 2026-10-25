"""Exports: the list as shown, as Excel and as CSV in the UI language."""

import io
from datetime import UTC, datetime, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from app.modules.plugins.contract import Kind, Movement

T = datetime(2025, 1, 10, 12, tzinfo=UTC)


def _load(authed, db):
    from app.modules.raw import service as raw
    from app.modules.sources.models import Source

    authed.post("/api/sources", json={"provider": "manual", "label": "Börse"})
    src = db.query(Source).one()
    raw.import_movements(
        db,
        src,
        "file",
        "t",
        [
            Movement(
                timestamp=T,
                kind=Kind.TRADE,
                incoming_asset="BTC",
                incoming_amount=Decimal("0.5"),
                outgoing_asset="EUR",
                outgoing_amount=Decimal("1234.5"),
            ),
            Movement(
                timestamp=T + timedelta(days=1),
                kind=Kind.SEND,
                outgoing_asset="BTC",
                outgoing_amount=Decimal("0.1"),
            ),
        ],
    )


def test_transactions_as_csv_in_german_and_as_excel(authed, db):
    _load(authed, db)
    r = authed.get("/api/export/transactions", params={"format": "csv", "tz": "Europe/Berlin"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    lines = text.strip().split("\r\n")
    assert lines[0].startswith("Datum;Art;Quelle") and len(lines) == 3
    assert "1234,5;EUR" in lines[2] and lines[2].startswith("2025-01-10 13:00:00;Handel;Börse")
    assert "kein Gegenstück" in lines[1]  # newest first, like the list
    only_sends = authed.get(
        "/api/export/transactions", params={"format": "csv", "kind": "send"}
    ).content
    assert only_sends.decode("utf-8-sig").count("\r\n") == 2
    wb = load_workbook(io.BytesIO(authed.get("/api/export/transactions").content))
    ws = wb.active
    assert ws.title == "Transaktionen" and ws.max_row == 3 and ws["E3"].value == 0.5


def test_holdings_export(authed, db):
    _load(authed, db)
    wb = load_workbook(io.BytesIO(authed.get("/api/export/holdings").content))
    ws = wb.active
    assert ws["A1"].value == "Coin" and ws["A2"].value == "BTC" and ws["C2"].value == 0.4


def test_a_futures_position_is_valued_once_in_euros_and_keeps_its_sign(db):
    """Review 2026-10-24: the display-currency value was converted a second time, a loss lost
    its minus."""
    from decimal import Decimal

    from app.modules.merged import service as merged

    position = {"type": "position", "result_value": "-108.00", "result_eur": "-100.00"}
    assert merged.item_values(db, None, [position]) == {id(position): Decimal("-100.00")}
