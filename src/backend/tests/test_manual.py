"""Manual entries (ADR 0036): exchange without connection, form, edit/delete with history,
template."""

import io
from datetime import UTC, datetime

from openpyxl import Workbook, load_workbook

TRADE = {
    "kind": "trade",
    "time": "2024-03-15T14:05",
    "incoming_asset": "eth",
    "incoming_amount": "0.5",
    "outgoing_asset": "USDT",
    "outgoing_amount": "1.500,25",
    "fee_asset": "USDT",
    "fee_amount": "1.5",
    "comment": "Gate.io",
}


def manual_exchange(authed, name="Gate.io") -> int:
    r = authed.post("/api/sources", json={"provider": "manual", "label": name})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_exchange_without_connection_is_offered_and_named_freely(authed):
    providers = {p["provider"]: p for p in authed.get("/api/plugins").json()["providers"]}
    assert providers["manual"]["kind"] == "exchange"
    assert [c["channel"] for c in providers["manual"]["channels"]] == ["manual"]
    a, b = manual_exchange(authed, "Gate.io"), manual_exchange(authed, "SafeTrade")
    labels = {s["id"]: s["label"] for s in authed.get("/api/sources").json()}
    assert labels[a] == "Gate.io" and labels[b] == "SafeTrade"


def test_form_entry_in_local_time_stored_in_utc(authed):
    sid = manual_exchange(authed)
    r = authed.post(f"/api/sources/{sid}/manual", json=TRADE)
    assert r.status_code == 201, r.text
    e = r.json()
    assert e["timestamp"].startswith("2024-03-15T13:05:00")  # Europe/Berlin, winter time
    assert e["time"] == "2024-03-15T14:05:00"
    assert (e["incoming_asset"], e["incoming_amount"]) == ("ETH", "0.5")
    assert e["outgoing_amount"] == "1500.25"
    listed = authed.get(f"/api/sources/{sid}/manual").json()
    assert listed["timezone"] == "Europe/Berlin" and [x["id"] for x in listed["entries"]] == [
        e["id"]
    ]
    assert authed.get("/api/raw", params={"source_id": sid}).json()["total"] == 1


def test_two_equal_form_entries_are_two(authed):
    sid = manual_exchange(authed)
    for _ in range(2):
        assert authed.post(f"/api/sources/{sid}/manual", json=TRADE).status_code == 201
    assert len(authed.get(f"/api/sources/{sid}/manual").json()["entries"]) == 2


def test_edit_keeps_the_old_row_and_only_the_new_counts(authed, db):
    from sqlalchemy import func, select

    from app.modules.raw.models import RawRetraction, RawTransaction

    sid = manual_exchange(authed)
    old = authed.post(f"/api/sources/{sid}/manual", json=TRADE).json()
    new = authed.put(f"/api/manual/{old['id']}", json={**TRADE, "incoming_amount": "0.6"}).json()
    assert new["id"] != old["id"] and new["edited"] and new["incoming_amount"] == "0.6"
    entries = authed.get(f"/api/sources/{sid}/manual").json()["entries"]
    assert [e["id"] for e in entries] == [new["id"]]
    assert authed.get("/api/raw", params={"source_id": sid}).json()["total"] == 1
    assert db.scalar(select(func.count()).select_from(RawTransaction)) == 2  # nothing lost
    assert db.get(RawRetraction, old["id"]).replaced_by_id == new["id"]
    again = authed.put(f"/api/manual/{old['id']}", json=TRADE)
    assert again.status_code == 409 and again.json()["error"]["code"] == "manual.already_changed"


def test_delete_hides_the_entry_everywhere(authed):
    sid = manual_exchange(authed)
    e = authed.post(f"/api/sources/{sid}/manual", json=TRADE).json()
    assert authed.delete(f"/api/manual/{e['id']}").status_code == 204
    assert authed.get(f"/api/sources/{sid}/manual").json()["entries"] == []
    assert authed.get("/api/raw", params={"source_id": sid}).json()["total"] == 0
    source = next(s for s in authed.get("/api/sources").json() if s["id"] == sid)
    assert source["rows"] == 0 and source["runs"] == []


def test_only_manual_rows_can_be_changed_and_the_manual_run_is_not_revertible(authed, db):
    from sqlalchemy import select

    from app.modules.raw.models import ImportRun

    sid = manual_exchange(authed)
    authed.post(f"/api/sources/{sid}/manual", json=TRADE)
    run = db.scalar(select(ImportRun).where(ImportRun.source_id == sid))
    r = authed.delete(f"/api/import-runs/{run.id}")
    assert r.status_code == 409 and r.json()["error"]["code"] == "imports.manual_run"
    assert authed.delete("/api/manual/999999").status_code == 404


def test_sides_are_checked_per_kind(authed):
    sid = manual_exchange(authed)

    def code(**changes):
        r = authed.post(f"/api/sources/{sid}/manual", json={**TRADE, **changes})
        return r.json()["error"]["code"], r.json()["error"]["params"]

    assert code(kind="receive") == ("manual.side_not_allowed", {"side": "outgoing"})
    assert code(outgoing_asset="", outgoing_amount="") == (
        "manual.side_missing",
        {"side": "outgoing"},
    )
    assert code(outgoing_amount="") == ("manual.side_incomplete", {"side": "outgoing"})
    assert code(incoming_amount="abc") == ("manual.invalid_amount", {"side": "incoming"})
    assert code(incoming_amount="-1") == ("manual.invalid_amount", {"side": "incoming"})
    assert code(kind="income", outgoing_asset="", outgoing_amount="", fee_asset="", fee_amount="")[
        0
    ] == ("manual.invalid_income_kind")
    assert code(time="15.03.2024")[0] == "manual.invalid_time"
    ok = authed.post(
        f"/api/sources/{sid}/manual",
        json={
            "kind": "income",
            "income_kind": "staking",
            "time": "2024-07-01T00:00",
            "incoming_asset": "DOT",
            "incoming_amount": "0.1",
        },
    )
    assert ok.status_code == 201


def test_manual_entries_on_an_exchange_with_connection(authed):
    authed.post("/api/sources", json={"provider": "bitget", "label": "Bitget"})
    sid = next(s["id"] for s in authed.get("/api/sources").json() if s["provider"] == "bitget")
    card = {
        "kind": "trade",
        "time": "2024-02-13T21:20:51",
        "incoming_asset": "BNB",
        "incoming_amount": "0.5142",
        "outgoing_asset": "EUR",
        "outgoing_amount": "150",
    }
    assert authed.post(f"/api/sources/{sid}/manual", json=card).status_code == 201


def _upload(authed, sid, data: bytes, preview: bool):
    url = f"/api/sources/{sid}/manual/template" + ("/preview" if preview else "")
    return authed.post(url, files={"file": ("vorlage.xlsx", data, "application/octet-stream")})


def test_template_download_upload_and_second_upload_adds_nothing(authed):
    sid = manual_exchange(authed)
    r = authed.get("/api/manual/template", params={"lang": "de"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    wb = load_workbook(io.BytesIO(r.content))
    assert (
        wb.sheetnames == ["Vorgaenge", "Hilfe"]
        and wb.active["D2"].value == "Kauf/Verkauf"
        and wb.active["C1"].value == "Markt"
    )
    preview = _upload(authed, sid, r.content, preview=True).json()
    assert (preview["new"], preview["duplicates"], preview["rejected"]) == (1, 0, 0)
    assert preview["kinds"] == {"trade": 1}
    assert (
        authed.get(f"/api/sources/{sid}/manual").json()["entries"] == []
    )  # preview writes nothing
    assert _upload(authed, sid, r.content, preview=False).json()["added"] == 1
    second = _upload(authed, sid, r.content, preview=False).json()
    assert (second["added"], second["duplicates"]) == (0, 1)


def _sheet(
    rows,
    headings=(
        "Date",
        "Time",
        "Type",
        "Received amount",
        "Received asset",
        "Spent amount",
        "Spent asset",
        "Note",
    ),
) -> bytes:
    wb = Workbook()
    wb.active.append(list(headings))
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_template_english_headings_equal_rows_and_named_problems(authed):
    sid = manual_exchange(authed, "SafeTrade")
    buy = ["2024-05-02", "10:00", "Buy/Sell", "100", "SAFE", "5", "USDT", None]
    data = _sheet(
        [
            buy,
            buy,  # two real equal buys stay two
            ["02.05.2024", "11:30", "Crypto deposit", "50", "USDT", None, None, "from Bitget"],
            ["2024-05-03", "09:00", "Unknown", "1", "BTC", None, None, None],
            ["2024-05-03", "09:00", "Crypto withdrawal", "1", "BTC", None, None, None],
            [None] * 8,
        ]
    )
    result = _upload(authed, sid, data, preview=False).json()
    assert (result["added"], result["rejected"]) == (3, 2)
    assert result["problems"] == [
        {"row": 5, "code": "manual.invalid_kind", "params": {}},
        {"row": 6, "code": "manual.side_missing", "params": {"side": "outgoing"}},
    ]
    entries = authed.get(f"/api/sources/{sid}/manual").json()["entries"]
    assert sorted(e["kind"] for e in entries) == ["receive", "trade", "trade"]


def test_template_without_required_columns_is_refused(authed):
    sid = manual_exchange(authed)
    r = _upload(authed, sid, _sheet([["x"]], headings=("Foo",)), preview=True)
    assert r.status_code == 400 and r.json()["error"]["code"] == "manual.template_columns"
    r = _upload(authed, sid, b"not an xlsx", preview=True)
    assert r.json()["error"]["code"] == "manual.template_unreadable"


def test_backup_export_marks_retracted_rows(authed, db, tmp_path):
    from app.modules.backup import service as backup

    sid = manual_exchange(authed)
    old = authed.post(f"/api/sources/{sid}/manual", json=TRADE).json()
    authed.delete(f"/api/manual/{old['id']}")
    target = tmp_path / "x.enc"
    assert backup.write_export(db, target) == 1
    (rec,) = backup.read_export(target)
    assert rec["retracted"] == {"replaced_by_fingerprint": None}


def test_exchange_with_only_deleted_entries_can_be_removed(authed):
    sid = manual_exchange(authed)
    e = authed.post(f"/api/sources/{sid}/manual", json=TRADE).json()
    edited = authed.put(f"/api/manual/{e['id']}", json=TRADE).json()
    assert authed.delete(f"/api/sources/{sid}").status_code == 409  # a live entry is left
    authed.delete(f"/api/manual/{edited['id']}")
    assert authed.delete(f"/api/sources/{sid}").status_code == 204
    assert all(s["id"] != sid for s in authed.get("/api/sources").json())


def test_seconds_are_kept(authed):
    sid = manual_exchange(authed)
    e = authed.post(
        f"/api/sources/{sid}/manual", json={**TRADE, "time": "2024-03-15T14:05:37"}
    ).json()
    assert e["time"] == "2024-03-15T14:05:37" and e["timestamp"].startswith("2024-03-15T13:05:37")


FUT = {
    "kind": "realized_pnl",
    "market": "futures",
    "time": "2024-06-01T10:00:05",
    "incoming_asset": "USDT",
    "incoming_amount": "12,5",
    "instrument": "btcusdt",
}


def test_futures_by_hand_one_sided_with_contract(authed):
    sid = manual_exchange(authed)
    e = authed.post(f"/api/sources/{sid}/manual", json=FUT).json()
    assert (e["market"], e["kind"], e["instrument"], e["incoming_amount"]) == (
        "futures",
        "realized_pnl",
        "BTCUSDT",
        "12.5",
    )
    loss = {
        **FUT,
        "incoming_asset": "",
        "incoming_amount": "",
        "outgoing_asset": "USDT",
        "outgoing_amount": "3,2",
        "fee_asset": "USDT",
        "fee_amount": "0,1",
    }
    assert authed.post(f"/api/sources/{sid}/manual", json=loss).status_code == 201
    rows = authed.get("/api/raw", params={"source_id": sid, "market": "futures"}).json()["rows"]
    assert len(rows) == 2

    def code(**changes):
        return authed.post(f"/api/sources/{sid}/manual", json={**FUT, **changes}).json()["error"][
            "code"
        ]

    assert code(outgoing_asset="USDT", outgoing_amount="1") == "manual.one_side"  # gain AND loss
    assert code(kind="trade") == "manual.invalid_kind"  # no spot kinds on futures
    assert code(market="spot") == "manual.invalid_kind"  # no realized result on spot
    assert code(kind="funding", fee_asset="USDT", fee_amount="1") == "manual.side_not_allowed"


def test_spot_transfer_and_migration(authed):
    sid = manual_exchange(authed)
    out = {
        "kind": "internal_transfer",
        "time": "2024-06-01T09:00",
        "outgoing_asset": "USDT",
        "outgoing_amount": "100",
    }
    assert authed.post(f"/api/sources/{sid}/manual", json=out).status_code == 201
    empty = authed.post(
        f"/api/sources/{sid}/manual", json={"kind": "internal_transfer", "time": "2024-06-01T09:00"}
    )
    assert empty.json()["error"]["code"] == "manual.side_missing"
    mig = {
        "kind": "migration",
        "time": "2024-09-10T12:00",
        "incoming_asset": "POL",
        "incoming_amount": "10",
        "outgoing_asset": "MATIC",
        "outgoing_amount": "10",
    }
    assert authed.post(f"/api/sources/{sid}/manual", json=mig).status_code == 201


def test_template_futures_rows(authed):
    sid = manual_exchange(authed)
    data = _sheet(
        [
            [
                "2024-06-01",
                "10:00:05",
                "Futures",
                "Realisiertes Ergebnis",
                "12,5",
                "USDT",
                None,
                None,
                "ETHUSDT",
            ],
            ["2024-06-01", "11:00", "Spot", "Kauf/Verkauf", "1", "ETH", "3000", "USDT", None],
            ["2024-06-01", "12:00", "Mond", "Kauf/Verkauf", "1", "ETH", "3000", "USDT", None],
        ],
        headings=(
            "Datum",
            "Uhrzeit",
            "Markt",
            "Art",
            "Erhalten Menge",
            "Erhalten Asset",
            "Abgegeben Menge",
            "Abgegeben Asset",
            "Kontrakt",
        ),
    )
    r = _upload(authed, sid, data, preview=False).json()
    assert (r["added"], r["rejected"]) == (2, 1) and r["problems"][0][
        "code"
    ] == "manual.invalid_market"
    fut = authed.get("/api/raw", params={"source_id": sid, "market": "futures"}).json()["rows"]
    assert fut[0]["instrument"] == "ETHUSDT"
    assert datetime.fromisoformat(fut[0]["timestamp"]) == datetime(2024, 6, 1, 8, 0, 5, tzinfo=UTC)


def test_a_wallet_gets_a_network_kept_by_hand(authed):
    """A chain without plugin (Optimism, Vanixxx 2026-10-11): named freely, address optional,
    never fetched, entries through the form."""
    providers = {p["provider"]: p for p in authed.get("/api/plugins").json()["providers"]}
    assert providers["manual_network"]["kind"] == "wallet"
    wallet = authed.post("/api/wallets", json={"name": "Coinbase Wallet"}).json()["id"]
    body = {"provider": "manual_network", "label": "Optimism", "wallet_id": wallet}
    assert authed.post("/api/sources", json={**body, "wallet_id": None}).status_code == 400
    r = authed.post("/api/sources", json={**body, "address": "0xabc"})
    assert r.status_code == 201, r.text
    sid = r.json()["id"]
    src = next(s for s in authed.get("/api/sources").json() if s["id"] == sid)
    assert (src["label"], src["wallet_id"], src["auto_fetch"]) == ("Optimism", wallet, False)
    assert authed.post("/api/sources", json=body).status_code == 201  # no address is fine
    assert authed.post(f"/api/sources/{sid}/manual", json=TRADE).status_code == 201
    assert authed.get("/api/raw", params={"source_id": sid}).json()["total"] == 1


def test_income_on_the_futures_account(authed):
    """A credit of the exchange to the futures account (Bitget risk capital, 2026-10-11)."""
    sid = manual_exchange(authed)
    body = {
        "kind": "income",
        "income_kind": "reward",
        "market": "futures",
        "time": "2024-04-12T20:35",
        "incoming_asset": "USDT",
        "incoming_amount": "11.2",
    }
    r = authed.post(f"/api/sources/{sid}/manual", json=body)
    assert r.status_code == 201, r.text
    assert (r.json()["market"], r.json()["kind"]) == ("futures", "income")


def test_the_hour_that_exists_twice_keeps_its_moment():
    """2026-10-25 02:30 Berlin exists twice (clocks go back): the reading matching the row's
    time so far wins (review 2026-10-24)."""
    from datetime import UTC, datetime

    from app.modules.manual.service import local_time

    second = datetime(2026, 10, 25, 1, 30, tzinfo=UTC)  # 02:30 CET, the second one
    got = local_time("2026-10-25T02:30:00", "Europe/Berlin", near=second)
    assert got == second
    first = local_time("2026-10-25T02:30:00", "Europe/Berlin")
    assert first.astimezone(UTC) == datetime(2026, 10, 25, 0, 30, tzinfo=UTC)  # 02:30 CEST


def test_numbers_from_excel_cells_lose_their_float_noise():
    """Review 2026-10-24: 0.1 + 0.2 came in as 0.30000000000000004."""
    from app.modules.manual.template import _cell_text

    assert _cell_text(0.1 + 0.2) == "0.3"
    assert _cell_text(1e-8) == "0.00000001"
    assert _cell_text(3.0) == "3" and _cell_text("1,5") == "1,5"
