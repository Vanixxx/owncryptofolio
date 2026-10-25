"""Tax pass (ADR 0042) on small, hand-checkable histories."""

from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import pytest

from app.modules.plugins.contract import IncomeKind, Kind, Market, Movement
from app.modules.tax.lots import Lot, Lots, Method, one_year_after, tax_free

D = Decimal
T = datetime(2024, 1, 10, 12, tzinfo=UTC)


def at(days: int, minutes: int = 0) -> datetime:
    return T + timedelta(days=days, minutes=minutes)


@pytest.fixture
def setup(authed, db):
    from app.modules.prices.models import PriceDaily
    from app.modules.raw import service as raw
    from app.modules.sources.models import Source

    authed.post("/api/sources", json={"provider": "manual", "label": "Exchange"})
    authed.post("/api/sources", json={"provider": "manual", "label": "Wallet"})
    a, b = db.query(Source).order_by(Source.id).all()
    # BTC 30,000 € on day 0, then 40,000 €; ETH 2,000 €; BNB 500 €; USDT 1 €
    for i in range(800):
        day = date(2024, 1, 10) + timedelta(days=i)
        for asset, close in (
            ("BTC", 30000 if i == 0 else 40000),
            ("ETH", 2000),
            ("BNB", 500),
            ("USDT", 1),
        ):
            db.add(PriceDaily(asset=asset, day=day, quote="EUR", close=D(close), source="t"))
    db.commit()

    def load(source, *moves):
        raw.import_movements(db, source, "file", "t", list(moves))

    return a, b, load


def run(db, method=Method.FIFO, **kw):
    from app.modules.merged import service as merged
    from app.modules.prices.service import Valuer
    from app.modules.tax.engine import Engine, Params

    return Engine(merged.build(db), Valuer(db), Params(method=method, **kw)).run()


def buy(when, qty, eur, asset="BTC", **kw):
    return Movement(
        timestamp=when,
        kind=Kind.TRADE,
        incoming_asset=asset,
        incoming_amount=D(qty),
        outgoing_asset="EUR",
        outgoing_amount=D(eur),
        **kw,
    )


def sell(when, qty, eur, asset="BTC", **kw):
    return Movement(
        timestamp=when,
        kind=Kind.TRADE,
        incoming_asset="EUR",
        incoming_amount=D(eur),
        outgoing_asset=asset,
        outgoing_amount=D(qty),
        **kw,
    )


# --- lots ----------------------------------------------------------------------------------


def test_holding_period_ends_with_the_anniversary():
    bought = datetime(2024, 3, 1, 10, tzinfo=UTC)
    assert one_year_after(date(2024, 2, 29)) == date(2025, 2, 28)
    assert not tax_free(bought, datetime(2025, 3, 1, 20, tzinfo=UTC))
    assert tax_free(bought, datetime(2025, 3, 2, 8, tzinfo=UTC))
    # German day: 23:30 UTC on 1 March is already 2 March in Berlin
    assert tax_free(bought, datetime(2025, 3, 1, 23, 30, tzinfo=UTC))


def test_fifo_and_lifo_differ_only_in_order_even_within_one_second():
    for method, first in ((Method.FIFO, D(10)), (Method.LIFO, D(20))):
        lots = Lots(method)
        place = (1, "spot", "X")
        lots.add(place, Lot("X", D(1), T, D(10)))
        lots.add(place, Lot("X", D(1), T, D(20)))  # same second
        (piece,), short = lots.take(place, D(1))
        assert piece.cost == first and short == 0


def test_network_fee_keeps_the_cost_with_the_rest():
    lots = Lots(Method.FIFO)
    place = (1, "spot", "X")
    lots.add(place, Lot("X", D(10), T, D(100)))
    assert lots.shrink(place, D(1)) == 0
    ((_, lot),) = lots.open()
    assert lot.qty == 9 and lot.cost == 100


# --- the pass ------------------------------------------------------------------------------


def test_sale_splits_over_lots_with_their_own_holding_period(setup, db):
    a, _, load = setup
    load(
        a,
        buy(at(0), 1, 30000),
        buy(at(300), 1, 40000),
        sell(at(400), "1.5", 60000),
    )
    fifo = run(db).disposals
    assert [(d.qty, d.cost, d.proceeds, d.tax_free) for d in fifo] == [
        (1, 30000, 40000, True),
        (D("0.5"), 20000, 20000, False),
    ]
    lifo = run(db, Method.LIFO).disposals
    assert [(d.qty, d.cost, d.tax_free) for d in lifo] == [
        (1, 40000, False),
        (D("0.5"), 15000, True),
    ]


def test_fees_as_the_bmf_reads_them(setup, db):
    a, _, load = setup
    load(
        a,
        # buy for euro, fee in euro: adds to the cost
        buy(at(1), 1, 40000, fee_asset="EUR", fee_amount=D(100)),
        # swap BTC → ETH, fee in the coin bought (Bitget): lowers proceeds and the new cost
        Movement(
            timestamp=at(2),
            kind=Kind.TRADE,
            outgoing_asset="BTC",
            outgoing_amount=D("0.5"),
            incoming_asset="ETH",
            incoming_amount=D(10),
            fee_asset="ETH",
            fee_amount=D("0.1"),
        ),
        # sell ETH for euro, fee in BNB (a third coin): BNB sold too, proceeds lower
        Movement(timestamp=at(3), kind=Kind.RECEIVE, incoming_asset="BNB", incoming_amount=D(1)),
        sell(at(4), 5, 10000, asset="ETH", fee_asset="BNB", fee_amount=D("0.1")),
    )
    res = run(db)
    swap, bnb, eth = res.disposals
    assert (swap.asset, swap.cost, swap.proceeds) == ("BTC", D(20050), D(19800))
    assert (bnb.asset, bnb.qty, bnb.proceeds) == ("BNB", D("0.1"), D(50))
    assert (eth.asset, eth.qty, eth.proceeds) == ("ETH", 5, D(9950))
    # 9.9 ETH cost 19,800 → 5 ETH cost 10,000
    assert eth.cost == D(10000)


def test_transfer_keeps_the_day_and_the_cost(setup, db):
    a, b, load = setup
    load(a, buy(at(0), 1, 30000))
    load(
        a,
        Movement(
            timestamp=at(10),
            kind=Kind.SEND,
            outgoing_asset="BTC",
            outgoing_amount=D(1),
            external_id="tx1",
        ),
    )
    load(
        b,
        Movement(
            timestamp=at(10, 5),
            kind=Kind.RECEIVE,
            incoming_asset="BTC",
            incoming_amount=D("0.999"),
            external_id="tx1",
        ),
        sell(at(400), "0.999", 50000),
    )
    (d,) = run(db).disposals
    assert d.source == b.id and d.cost == 30000 and d.tax_free and d.acquired == at(0)


def test_migration_is_no_sale_and_keeps_the_day(setup, db):
    a, _, load = setup
    load(
        a,
        buy(at(0), 1, 2000, asset="NGL"),
        Movement(
            timestamp=at(100),
            kind=Kind.MIGRATION,
            outgoing_asset="NGL",
            outgoing_amount=D(1),
            incoming_asset="NTGL",
            incoming_amount=D(1000),
        ),
    )
    res = run(db)
    assert res.disposals == []
    ((place, lot),) = res.lots.open()
    assert (place[2], lot.qty, lot.cost, lot.acquired) == ("NTGL", 1000, 2000, at(0))


def test_marks_external_lost_and_own_wallet(setup, db):
    from app.modules.merged import service as merged
    from app.modules.raw.models import RawTransaction

    a, _, load = setup
    load(
        a,
        buy(at(0), 3, 90000),
        *(
            Movement(
                timestamp=at(5, i),
                kind=Kind.SEND,
                outgoing_asset="BTC",
                outgoing_amount=D(1),
                counterparty=f"addr{i}",
            )
            for i in range(3)
        ),
    )
    sends = db.query(RawTransaction).filter_by(kind="send").order_by(RawTransaction.id).all()
    for r, mark in zip(sends, ("external", "lost", "own_untracked"), strict=True):
        merged.mark_open(db, r.id, mark, None)
    res = run(db)
    (paid,) = res.disposals
    assert (paid.proceeds, paid.cost, paid.how) == (40000, 30000, "external")
    (lost,) = res.losses
    assert (lost.qty, lost.cost, lost.acquired) == (1, 30000, at(0))
    ((place, lot),) = res.lots.open()
    assert place[0] == 0 and lot.acquired == at(0)


def test_income_is_taxed_at_inflow_and_starts_its_own_lot(setup, db):
    a, _, load = setup
    load(
        a,
        Movement(
            timestamp=at(0),
            kind=Kind.INCOME,
            income_kind=IncomeKind.STAKING,
            incoming_asset="ETH",
            incoming_amount=D("0.5"),
        ),
        Movement(
            timestamp=at(1),
            kind=Kind.INCOME,
            income_kind=IncomeKind.AIRDROP,
            incoming_asset="ETH",
            incoming_amount=D(1),
        ),
    )
    res = run(db)
    (inc,) = res.incomes
    assert (inc.value, inc.kind) == (1000, "staking")
    assert sorted(lot.cost for _, lot in res.lots.open()) == [0, 1000]
    taxed = run(db, airdrop_taxable=True)
    assert len(taxed.incomes) == 2


def test_sale_without_lot_is_uncovered_without_cost(setup, db):
    a, _, load = setup
    load(a, buy(at(0), 1, 30000), sell(at(1), 2, 80000))
    res = run(db)
    (d,) = res.disposals
    (u,) = res.uncovered
    assert (d.qty, d.proceeds) == (1, 40000)
    assert (u.qty, u.proceeds, u.reason) == (1, 40000, "no_lot")


def test_futures_flows_are_listed_and_move_the_margin_coin(setup, db):
    a, _, load = setup
    load(
        a,
        Movement(timestamp=at(0), kind=Kind.RECEIVE, incoming_asset="USDT", incoming_amount=D(100)),
        Movement(
            timestamp=at(1),
            kind=Kind.REALIZED_PNL,
            market=Market.FUTURES,
            incoming_asset="USDT",
            incoming_amount=D(30),
            fee_asset="USDT",
            fee_amount=D(1),
        ),
        Movement(
            timestamp=at(2),
            kind=Kind.FUNDING,
            market=Market.FUTURES,
            outgoing_asset="USDT",
            outgoing_amount=D(2),
        ),
    )
    res = run(db)
    assert [(f.kind, f.value) for f in res.futures] == [("pnl", 30), ("fee", -1), ("funding", -2)]
    # spent by the futures account: out of the lots, but no sale under § 23 (like Blockpit)
    assert res.disposals == []
    assert sum(lot.qty for _, lot in res.lots.open()) == 100 + 30 - 1 - 2


def test_futures_bonus_is_section_20_not_income(setup, db):
    a, _, load = setup
    load(
        a,
        Movement(
            timestamp=at(0),
            kind=Kind.INCOME,
            income_kind=IncomeKind.REWARD,
            market=Market.FUTURES,
            incoming_asset="USDT",
            incoming_amount=D(20),
            comment="BONUS",
        ),
        Movement(
            timestamp=at(5),
            kind=Kind.FEE,
            market=Market.FUTURES,
            fee_asset="USDT",
            fee_amount=D(20),
            comment="BONUS",
        ),
    )
    res = run(db)
    assert res.incomes == []
    assert [(f.kind, f.value) for f in res.futures] == [("bonus", 20), ("bonus", -20)]


def test_card_purchase_without_euro_is_estimated(setup, db):
    a, _, load = setup
    load(
        a,
        Movement(
            timestamp=at(0),
            kind=Kind.OTHER,
            incoming_asset="BTC",
            incoming_amount=D(1),
            comment="Buy with card",
        ),
    )
    ((_, lot),) = run(db).lots.open()
    assert lot.estimated and lot.cost == 30000


# --- report, keeping, files ----------------------------------------------------------------


def test_year_report_limit_and_freeze_notices_a_later_change(setup, authed, db):
    a, _, load = setup
    load(a, buy(at(0), 1, 30000), sell(at(10), "0.1", 4000))
    body = authed.get("/api/tax/2024").json()
    s23 = body["summary"]["s23"]
    # 0.1 BTC: proceeds 4,000, cost 3,000 → 1,000 € gain: reaches the 1,000 € limit of 2024
    assert (s23["net"], s23["limit"], s23["taxable"]) == ("1000.00", "1000.00", True)
    assert authed.get("/api/tax/years").json() == [2024]
    c = body["summary"]["carry"]["year"]
    assert (c["net"], c["taxable"], c["end"]) == ("1000.00", True, "0.00")

    assert authed.post("/api/tax/2024/freeze").status_code == 200
    assert [f["same"] for f in authed.get("/api/tax/2024").json()["frozen"]] == [True]
    # a correction (here: one more sale) changes the year → the kept report says so
    load(a, sell(at(20), "0.1", 4000))
    assert [f["same"] for f in authed.get("/api/tax/2024").json()["frozen"]] == [False]

    lifo = authed.get("/api/tax/2024?method=lifo").json()["summary"]
    assert lifo["method"] == "lifo"


def test_files_and_holding_periods(setup, authed, db):
    a, _, load = setup
    load(a, buy(at(0), 1, 30000), buy(datetime.now(UTC) - timedelta(days=10), 1, 40000))
    x = authed.get("/api/tax/2024/file?kind=xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    p = authed.get("/api/tax/2024/file?kind=pdf")
    assert p.status_code == 200 and p.content[:4] == b"%PDF"
    h = authed.get("/api/tax/holdings").json()["BTC"]
    assert h["qty"] == "2" and h["free"] == "1"
    f = authed.get("/api/tax/free").json()
    btc = next(c for c in f["coins"] if c["asset"] == "BTC")
    assert (btc["free"], btc["taxed"], btc["next_qty"]) == ("1", "1", "1")
    assert btc["next_free"] is not None and set(f["total"]) == {
        "value_free",
        "value_taxed",
        "gain_free",
        "gain_taxed",
    }


def test_rebooking_booked_arrival_first_gets_the_real_lots(setup, db):
    """Kraken staking and Bitget futures book the arrival before the departure."""
    a, _, load = setup
    load(
        a,
        buy(at(0), 1, 30000),
        Movement(
            timestamp=at(5),
            kind=Kind.INTERNAL_TRANSFER,
            market=Market.FUTURES,
            incoming_asset="BTC",
            incoming_amount=D(1),
        ),
        Movement(
            timestamp=at(5, 1),
            kind=Kind.INTERNAL_TRANSFER,
            outgoing_asset="BTC",
            outgoing_amount=D(1),
        ),
    )
    res = run(db)
    ((place, lot),) = res.lots.open()
    assert place[1] == "futures" and lot.acquired == at(0) and lot.cost == 30000
    assert not lot.estimated and res.uncovered == []


def test_compare_kept_reports_line_by_line(setup, authed, db):
    a, _, load = setup
    load(a, buy(at(0), 1, 30000), sell(at(10), "0.1", 4000))
    first = authed.post("/api/tax/2024/freeze", json={"note": "an den Steuerberater"}).json()["id"]
    load(a, sell(at(20), "0.1", 4500))
    second = authed.post("/api/tax/2024/freeze").json()["id"]
    kept = authed.get("/api/tax/2024").json()["frozen"]
    assert [(k["id"], k["note"]) for k in kept] == [(second, None), (first, "an den Steuerberater")]

    c = authed.get(f"/api/tax/2024/compare?a={first}&b={second}").json()
    net = next(f for f in c["figures"] if (f["pot"], f["key"]) == ("s23", "net"))
    assert (net["old"], net["new"], net["delta"]) == ("1000.00", "2500.00", "1500.00")
    sales = [x for x in c["changes"] if x["cat"] == "disposals"]
    assert [(x["state"], x["new"]["amount"]) for x in sales] == [("added", "1500.00")]
    # the year-end holding of BTC shrank by the second sale
    (held,) = [x for x in c["changes"] if x["cat"] == "year_end"]
    assert (held["old"]["qty"], held["new"]["qty"]) == ("0.9", "0.8")
    # against now with nothing changed since the second: no line differs
    assert authed.get(f"/api/tax/2024/compare?a={second}").json()["changes"] == []
    assert authed.get("/api/tax/2024/compare?a=999").status_code == 404


def test_tax_side_of_a_buy_a_sale_and_a_transfer(setup, authed, db):
    from app.modules.raw.models import RawTransaction

    a, b, load = setup
    load(a, buy(at(0), 1, 30000))
    load(
        a,
        Movement(
            timestamp=at(10),
            kind=Kind.SEND,
            outgoing_asset="BTC",
            outgoing_amount=D(1),
            external_id="tx9",
        ),
    )
    load(
        b,
        Movement(
            timestamp=at(10, 5),
            kind=Kind.RECEIVE,
            incoming_asset="BTC",
            incoming_amount=D(1),
            external_id="tx9",
        ),
        sell(at(400), "0.4", 20000),
    )
    ids = {r.kind: r.id for r in db.query(RawTransaction)}
    rows = db.query(RawTransaction).filter_by(kind="trade").order_by(RawTransaction.timestamp).all()
    bought, sold = rows[0].id, rows[1].id

    b_ = authed.get(f"/api/tax/rows?ids={bought}").json()
    assert [c["qty"] for c in b_["created"]] == ["1"]
    assert [(s["qty"], s["tax_free"]) for s in b_["later"]["sold"]] == [("0.4", True)]
    assert [o["qty"] for o in b_["later"]["open"]] == ["0.6"]

    s_ = authed.get(f"/api/tax/rows?ids={sold}").json()
    ((d),) = s_["sold"]
    assert (d["origin"], d["gain"]) == (bought, "8000.00")

    t_ = authed.get(f"/api/tax/rows?ids={ids['send']},{ids['receive']}").json()
    assert [(m["qty"], m["origin"]) for m in t_["moved"]] == [("1", bought)]
    assert t_["sold"] == []


def test_keeping_stores_excel_and_pdf_as_they_were(setup, authed, db):
    from app.modules.tax.models import TaxReport

    a, _, load = setup
    load(a, buy(at(0), 1, 30000), sell(at(10), "0.1", 4000))
    kept = authed.post("/api/tax/2024/freeze").json()["id"]
    pdf = authed.get(f"/api/tax/kept/{kept}/file?kind=pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    assert "festgehalten" in pdf.headers["content-disposition"]
    xlsx = authed.get(f"/api/tax/kept/{kept}/file?kind=xlsx")
    assert xlsx.content[:2] == b"PK"
    # later changes leave the kept document alone
    load(a, sell(at(20), "0.1", 4500))
    assert authed.get(f"/api/tax/kept/{kept}/file?kind=pdf").content == pdf.content
    # a report kept before files existed gets them from its kept lines
    row = db.get(TaxReport, kept)
    row.pdf = None
    db.commit()
    assert authed.get(f"/api/tax/kept/{kept}/file?kind=pdf").content[:4] == b"%PDF"
    assert authed.get("/api/tax/kept/999/file?kind=pdf").status_code == 404


def test_pdf_summary_fits_on_the_first_page(setup, authed, db):
    import re

    from app.modules.tax import report
    from app.modules.tax.files import pdf, tables

    a, _, load = setup
    load(a, buy(at(0), 1, 30000), sell(at(10), "0.1", 4000))
    rep = report.year_report(db, 2024, None)
    pages = len(re.findall(rb"/Type\s*/Page\b(?!s)", pdf(rep)))
    # one page of summary, then one landscape page per (short) detail table
    assert pages == 1 + len(tables(rep))

    # the longest summary there can be: every form line, every kind of income (and two the
    # app has no German name for), exchange bonus, KAP lines, a loss carry with a notice
    s = rep["summary"]
    s["forms"] = [
        {"form": f, "line": "1", "key": k, "value": "1"}
        for f, k in [
            ("SO", "s23"),
            ("SO", "s22"),
            ("KAP", "kap19"),
            ("KAP", "kap21"),
            ("KAP", "kap24"),
        ]
    ]
    s["s22"]["by_kind"] = {
        k: "1" for k in ["staking", "reward", "interest", "airdrop", "mining", "lending"]
    }
    s["s20"].update(bonus="1", kap_gains="1", kap_losses="1")
    keys = ["start", "used", "back_in", "back_out", "loss", "remaining", "end", "assessed"]
    s["carry"] = {"year": dict.fromkeys(keys, "1")}
    # and the longest head the settings allow
    s["holder"] = {"name": "N" * 120, "address": "Straße 1\n" + "A" * 290, "tax_number": "9" * 40}
    pages = len(re.findall(rb"/Type\s*/Page\b(?!s)", pdf(rep)))
    assert pages == 1 + len(tables(rep))


def test_tax_report_head_has_name_address_and_tax_number(setup, authed, db):
    import io

    from openpyxl import load_workbook

    a, _, load = setup
    load(a, buy(at(0), 1, 30000), sell(at(10), "0.1", 4000))
    # empty by default: no head lines
    book = load_workbook(io.BytesIO(authed.get("/api/tax/2024/file?kind=xlsx").content))
    assert book["Zusammenfassung"]["A1"].value != "Name"
    authed.patch(
        "/api/settings",
        json={
            "report.name": " Max Mustermann ",
            "report.address": "Musterstraße 1\n12345 Musterstadt",
            "report.tax_number": "123/456/78901",
        },
    )
    assert authed.get("/api/tax/2024").json()["summary"]["holder"] == {
        "name": "Max Mustermann",
        "address": "Musterstraße 1\n12345 Musterstadt",
        "tax_number": "123/456/78901",
    }
    book = load_workbook(io.BytesIO(authed.get("/api/tax/2024/file?kind=xlsx").content))
    rows = [[c.value for c in r] for r in book["Zusammenfassung"].iter_rows(max_row=3)]
    assert rows == [
        ["Name", "Max Mustermann"],
        ["Anschrift", "Musterstraße 1, 12345 Musterstadt"],
        ["Steuernummer", "123/456/78901"],
    ]
    # a kept report keeps the head it was made with
    kept = authed.post("/api/tax/2024/freeze").json()["id"]
    authed.patch("/api/settings", json={"report.tax_number": ""})
    old = load_workbook(io.BytesIO(authed.get(f"/api/tax/kept/{kept}/file?kind=xlsx").content))
    assert old["Zusammenfassung"]["A3"].value == "Steuernummer"
    assert authed.get("/api/tax/2024").json()["summary"]["holder"]["tax_number"] == ""


def test_pdf_stars_estimated_costs_instead_of_a_column():
    from app.modules.tax.files import _star_estimated

    cols = [("Coin", "text", 1), ("Kosten €", "eur", 1), ("Kosten geschätzt", "text", 1)]
    out_cols, rows, starred = _star_estimated(
        cols, [["BTC", D("12.3"), "ja"], ["ETH", D(5), "nein"]]
    )
    assert [c[0] for c in out_cols] == ["Coin", "Kosten €"] and starred
    assert rows == [["BTC", "12,30*"], ["ETH", D(5)]]
    assert _star_estimated(cols, [["ETH", D(5), "nein"]])[2] is False


def test_kept_version_note_can_change_and_version_can_go(setup, authed, db):
    a, _, load = setup
    load(a, buy(at(0), 1, 30000), sell(at(10), "0.1", 4000))
    kept = authed.post("/api/tax/2024/freeze").json()["id"]
    assert authed.patch(
        f"/api/tax/kept/{kept}", json={"note": "  an den Steuerberater "}
    ).json() == {
        "id": kept,
        "note": "an den Steuerberater",
    }
    assert authed.get("/api/tax/2024").json()["frozen"][0]["note"] == "an den Steuerberater"
    assert authed.delete(f"/api/tax/kept/{kept}").status_code == 204
    assert authed.get("/api/tax/2024").json()["frozen"] == []
    assert authed.delete(f"/api/tax/kept/{kept}").status_code == 404


def test_card_purchase_without_euro_is_a_notice(setup, authed, db):
    a, _, load = setup
    load(
        a,
        Movement(
            timestamp=at(0),
            kind=Kind.OTHER,
            incoming_asset="BTC",
            incoming_amount=D(1),
            comment="Buy with card",
        ),
    )
    body = authed.get("/api/merged?show=card").json()
    assert body["summary"]["card"] == 1 and len(body["items"]) == 1


def test_rounding_shortfall_is_not_an_open_point():
    lots = Lots(Method.FIFO)
    place = (1, "spot", "QUBIC")
    lots.add(place, Lot("QUBIC", D("11387279.7216"), T, D(100)))
    _, short = lots.take(place, D("11387280"))
    assert short == 0
    lots.add(place, Lot("QUBIC", D(1), T, D(1)))
    assert lots.take(place, D(2))[1] == 1


def test_upcoming_tax_free(setup, authed, db):
    from app.modules.prices.models import PriceLatest

    a, _, load = setup
    soon = datetime.now(UTC) - timedelta(days=355)  # free in about 11 days
    load(a, buy(soon, 1, 30000, asset="ETH"), buy(at(0), 1, 30000))
    db.add(PriceLatest(asset="ETH", quote="EUR", price=D(3000), source="t", at=datetime.now(UTC)))
    db.commit()
    (u,) = authed.get("/api/tax/upcoming?days=30").json()
    assert (u["asset"], u["qty"], u["value"], u["gain"]) == ("ETH", "1", "3000.00", "-27000.00")


def test_limits_of_this_year(setup, authed, db):
    from app.modules.prices.models import PriceDaily

    a, _, load = setup
    now = datetime.now(UTC)
    for i in range(3):
        db.add(
            PriceDaily(
                asset="ETH",
                day=(now - timedelta(days=i)).date(),
                quote="EUR",
                close=D(2000),
                source="t",
            )
        )
    db.commit()
    # bought for 1,150 €, sold for 2,000 €: 850 € gain = 85 % of 1,000 €
    load(
        a,
        buy(now - timedelta(days=2), 1, 1150, asset="ETH"),
        sell(now - timedelta(days=1), 1, 2000, asset="ETH"),
    )
    body = authed.get("/api/tax/limits").json()
    assert (body["s23"]["share"], body["s23"]["level"]) == ("85.00", "near")
    assert body["s22"]["level"] == "ok"


def test_a_year_with_only_uncovered_sales_is_listed(setup, authed, db):
    """Only the sales imported, not the buys: the year still shows, with the list that
    points at the missing purchases (review 2026-10-24)."""
    a, _, load = setup
    load(a, sell(at(10), "0.1", 4000))
    assert authed.get("/api/tax/years").json() == [2024]
    assert authed.get("/api/tax/2024").json()["summary"]["open"]["uncovered"] == 1
