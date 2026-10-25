"""Kraken — ledger export (research 2026-08-29).

* **Times are UTC**, stated nowhere: two withdrawals reappear as KuCoin deposits
  three minutes later under UTC, over an hour later under local time.
* `amount` excludes the fee, `fee` comes off additionally (checked against the
  `balance` column). A withdrawal's amount is what arrives.
* A trade is two rows sharing `refid`; its fee sits on exactly one of them. Fees
  on both would lose one and are rejected. `spend`/`receive` (instant buy) is a
  trade too.
* `DOT.S`, `ETH.F` are the same coin in another pot; old `XXBT`/`ZEUR` tickers
  are mapped (missing them would be silent: `ZEUR` would count as crypto).
* Spot↔futures transfers are booked as internal transfers but must net to zero
  per asset; a remainder would be a hidden futures result and rejects the file.
* Unknown booking types reject the file instead of being guessed.
* CSV is the better source: a workbook stores numbers as floats.
"""

from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from collections.abc import Iterator
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from app.modules.plugins import archive
from app.modules.plugins.contract import Kind, Movement

FIAT = {"EUR", "USD", "GBP", "CHF", "CAD", "AUD", "JPY"}
OLD = {
    "XXBT": "BTC",
    "XBT": "BTC",
    "XETH": "ETH",
    "XXRP": "XRP",
    "XLTC": "LTC",
    "XXLM": "XLM",
    "XXDG": "DOGE",
    "XXMR": "XMR",
    "XZEC": "ZEC",
    "XETC": "ETC",
    "XREP": "REP",
    "XMLN": "MLN",
    "ZEUR": "EUR",
    "ZUSD": "USD",
    "ZGBP": "GBP",
    "ZCAD": "CAD",
    "ZJPY": "JPY",
    "ZAUD": "AUD",
}
SUFFIX = re.compile(r"\.(S|M|F|B|P\d*|HOLD)$")
KNOWN = {"deposit", "withdrawal", "trade", "spend", "receive", "transfer"}
TRADES = {"trade", "spend", "receive"}
FUTURES_TRANSFER = {"spottofutures", "spotfromfutures"}


def _asset(raw) -> str:
    text = SUFFIX.sub("", str(raw or "").strip().upper())
    return OLD.get(text, text)


def _num(raw, where: str) -> Decimal:
    if raw is None or raw == "":
        return Decimal(0)
    try:
        return Decimal(str(raw).strip().replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"{where}: {raw!r} is not a number") from exc


def _time(raw, where: str) -> datetime:
    if isinstance(raw, datetime):
        return raw.replace(tzinfo=UTC)
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%dT%H:%M:%SZ"):
        try:
            return datetime.strptime(str(raw).strip(), fmt).replace(tzinfo=UTC)
        except ValueError:
            continue
    raise ValueError(f"{where}: unreadable time {raw!r}")


def _with_header(rows: list[tuple], where: str) -> list[dict]:
    start = next(
        (i for i, r in enumerate(rows) if r and str(r[0] or "").strip().lower() == "txid"), None
    )
    if start is None:
        raise ValueError(f"{where}: no header row starting with 'txid' (Kraken ledger export)")
    head = [str(c or "").strip() for c in rows[start]]
    return [
        dict(zip(head, r, strict=False))
        for r in rows[start + 1 :]
        if any(str(c or "").strip() for c in r)
    ]


def _rows(path: Path) -> list[dict]:
    files = (
        sorted(p for p in path.iterdir() if p.suffix.lower() in (".csv", ".xlsx"))
        if path.is_dir()
        else [path]
    )
    out: list[dict] = []
    for f in files:
        if f.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            archive.check(f)
            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb["ledgers"] if "ledgers" in wb.sheetnames else wb.worksheets[0]
            out += _with_header(list(ws.iter_rows(values_only=True)), f.name)
            wb.close()
        else:
            text = f.read_text(encoding="utf-8-sig", errors="replace")
            out += _with_header([tuple(r) for r in csv.reader(io.StringIO(text))], f.name)
    return out


def _raw(row: dict) -> dict:
    return {k: (v.isoformat() if isinstance(v, datetime) else v) for k, v in row.items() if k}


def read(path: Path, options: dict | None = None) -> Iterator[Movement]:
    rows = _rows(path)
    unknown = {str(r.get("type") or "").strip().lower() for r in rows} - KNOWN - {""}
    if unknown:
        raise ValueError(
            f"unknown booking types {sorted(unknown)} - decided on real data, not guessed"
        )
    rest: dict[str, Decimal] = defaultdict(Decimal)
    by_ref: dict[str, list[dict]] = defaultdict(list)
    out: list[Movement] = []
    for r in rows:
        typ = str(r.get("type") or "").strip().lower()
        sub = str(r.get("subtype") or "").strip().lower()
        where = f"{r.get('txid')} ({typ})"
        if typ == "transfer":
            if sub not in FUTURES_TRANSFER:
                raise ValueError(f"unknown transfer subtype {sub!r}")
            asset = _asset(r.get("asset"))
            amount = _num(r.get("amount"), where)
            rest[asset] += amount - _num(r.get("fee"), where)
            out.append(
                Movement(
                    timestamp=_time(r.get("time"), where),
                    kind=Kind.INTERNAL_TRANSFER,
                    **(
                        {"incoming_asset": asset, "incoming_amount": amount}
                        if amount > 0
                        else {"outgoing_asset": asset, "outgoing_amount": -amount}
                    ),
                    external_id=str(r.get("txid") or "") or None,
                    comment=sub,
                    raw=_raw(r),
                )
            )
        elif typ in ("deposit", "withdrawal"):
            out.append(_single(r, typ, where))
        elif typ in TRADES:
            by_ref[str(r.get("refid") or "")].append(r)
    left = {a: v for a, v in rest.items() if v != 0}
    if left:
        raise ValueError(
            f"spot-futures transfers do not net to zero {left}: a futures result is hidden in them"
        )
    yield from out
    for ref, group in by_ref.items():
        yield _trade(ref, group)


def _single(r: dict, typ: str, where: str) -> Movement:
    asset = _asset(r.get("asset"))
    amount, fee = _num(r.get("amount"), where), _num(r.get("fee"), where)
    fiat = asset in FIAT
    base = {
        "timestamp": _time(r.get("time"), where),
        "external_id": str(r.get("txid") or "") or None,
        "fee_asset": asset if fee else None,
        "fee_amount": fee or None,
        "raw": _raw(r),
    }
    if typ == "deposit":
        return Movement(
            kind=Kind.FIAT_DEPOSIT if fiat else Kind.RECEIVE,
            incoming_asset=asset,
            incoming_amount=amount,
            **base,
        )
    return Movement(
        kind=Kind.FIAT_WITHDRAWAL if fiat else Kind.SEND,
        outgoing_asset=asset,
        outgoing_amount=abs(amount),
        **base,
    )


def _trade(ref: str, rows: list[dict]) -> Movement:
    if len(rows) != 2:
        raise ValueError(f"refid {ref}: {len(rows)} rows instead of two")
    out = [r for r in rows if _num(r.get("amount"), ref) < 0]
    inc = [r for r in rows if _num(r.get("amount"), ref) > 0]
    if len(out) != 1 or len(inc) != 1:
        raise ValueError(f"refid {ref}: expected one negative and one positive row")
    fees = [
        (_asset(r.get("asset")), _num(r.get("fee"), ref)) for r in rows if _num(r.get("fee"), ref)
    ]
    if len(fees) > 1:
        raise ValueError(f"refid {ref}: fees on both sides {fees}")
    o, i = out[0], inc[0]
    return Movement(
        timestamp=_time(o.get("time"), ref),
        kind=Kind.TRADE,
        outgoing_asset=_asset(o.get("asset")),
        outgoing_amount=abs(_num(o.get("amount"), ref)),
        incoming_asset=_asset(i.get("asset")),
        incoming_amount=_num(i.get("amount"), ref),
        fee_asset=fees[0][0] if fees else None,
        fee_amount=fees[0][1] if fees else None,
        external_id=ref,
        raw={"refid": ref, "rows": [_raw(r) for r in rows]},
    )
