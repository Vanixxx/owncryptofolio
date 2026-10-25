"""MEXC — Excel exports (research 2026-08-24, re-checked 2026-10-03).

* **The time zone depends on the download moment**: the same withdrawal reads
  21:44:44 in a January download and 22:44:44 in an August one; true is 20:44:44
  UTC. Some columns say `(UTC+02:00)`, others nothing. Per file: the label in its
  own headers, else the download moment in its name, else a labelled sibling of
  the same export (same date-range suffix).
* **English and German exports** differ only in column names; files are
  recognised by their columns, not their names.
* The fee sticks to the amount (`0.04USDT`); without a currency (German export)
  it is the quote currency (0.1 % of the total, checked on real rows).
* Withdrawals: requested − fee = settled; the settled amount leaves, the fee
  separately.
* Transfers FIAT↔SPOT stay inside the spot side (net zero, shown); SPOT→FUTURES
  leaves spot, the futures side comes from the capital flow.
* **Futures bonus** (trial funds, re-checked 2026-10-11 with exports back to 2024-06):
  MEXC pays losses, fees and funding out of a bonus first. The capital flow then shows
  the loss *and* a `BONUS_DEDUCT` of the same size; the statement shows only the loss.
  So the bonus is income (`BONUS` inflow), an unused bonus that runs out leaves again
  (`BONUS` outflow, booked as a futures fee), and `BONUS_DEDUCT` is skipped. Measured:
  own money on futures 2024-10-01 = 53.53 USDT (the transfer of 2024-09-10),
  2026-08-24 = 87.98, as the snapshot says.
* Overlapping exports are fine: rows are converted to UTC, so the same row in two
  files has the same fingerprint and is recognised as a duplicate.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from app.modules.plugins import archive
from app.modules.plugins.contract import IncomeKind, Kind, Market, Movement

TZ_LABEL = re.compile(r"\(UTC([+-])(\d{1,2})(?::(\d{2}))?\)")
AMOUNT_UNIT = re.compile(r"^(-?[\d.]+)\s*([A-Z0-9]*)$")
MS_STAMP = re.compile(r"_(1\d{12})\b")
DATE_STAMP = re.compile(r"(\d{4}-\d{2}-\d{2})[ _]\d{2}[_:]\d{2}[_:]\d{2}")
RANGE = re.compile(r"(\d{4}-\d{2}-\d{2}-\d{4}-\d{2}-\d{2}|\d{8}-\d{8})")

COLS = {  # English -> German column names
    "Pairs": "Paare",
    "Time": "Zeit",
    "Side": "Seite",
    "Executed Amount": "Ausgeführter Betrag",
    "Total": "Gesamt",
    "Fee": "Gebühr",
    "Role": "Rolle",
    "Status": "Status",
    "Crypto": "Krypto",
    "Network": "Netzwerk",
    "Deposit Amount": "Einzahlungsbetrag",
    "Request Amount": "Angeforderter Betrag",
    "Settlement Amount": "Abrechnungsbetrag",
    "Trading Fee": "Handelsgebühr",
    "Withdrawal Address": "Auszahlungsadresse",
    "TxID": "TxID",
    # futures capital flow (Futures-Mittelfluss) and position history, German exports
    "Futures Trading Pair": "Futures-Handelspaar",
    "Fund Type": "Mittelart",
    "Fund Flow Type": "Art des Mittelflusses",
    "Amount": "Betrag",
    "Futures": "Handelspaar",
    "Open Time": "Öffnungszeiten",
    "Close Time": "Schließzeiten",
    "Margin Mode": "Margin-Modus",
    "Direction": "Richtung",
    "Realized PNL": "Realisierte PNL",
}
ROLES = {
    "trades": {"Pairs", "Side", "Executed Amount", "Total"},
    "deposits": {"Deposit Amount", "Crypto"},
    "withdrawals": {"Request Amount", "Settlement Amount"},
    "otc": {"Trading Token", "Settlement Token", "Order Amount"},
    "transfers": {"From System", "To System", "Currency"},
    "futures_flow": {"Fund Type", "Fund Flow Type", "Amount"},
    "positions": {"Futures", "Realized PNL", "Direction"},
}
SUCCESS = ("success", "erfolg", "complete", "credited")


def _table(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    archive.check(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [
        ["" if c is None else str(c) for c in r]
        for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)
    ]
    wb.close()
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], []
    head = [c.strip() for c in rows[0]]
    return head, [dict(zip(head, r, strict=False)) for r in rows[1:]]


def _english(head: list[str], rows: list[dict]) -> tuple[set[str], list[dict]]:
    """German column names to English; the time-zone label is cut off for matching."""
    back = {v: k for k, v in COLS.items()}
    names = {back.get(TZ_LABEL.sub("", h).strip(), TZ_LABEL.sub("", h).strip()): h for h in head}
    return set(names), [{k: r.get(v, "") for k, v in names.items()} for r in rows]


def _role(cols: set[str]) -> str | None:
    return next((role for role, need in ROLES.items() if need <= cols), None)


def _label_offset(head: list[str]) -> timedelta | None:
    for h in head:
        m = TZ_LABEL.search(h)
        if m and "Date" not in h:  # the daily snapshot column says UTC+8 and means a date
            sign, hh, mm = m.groups()
            delta = timedelta(hours=int(hh), minutes=int(mm or 0))
            return -delta if sign == "-" else delta
    return None


def _name_offset(name: str, tz: str) -> timedelta | None:
    if m := MS_STAMP.search(name):
        moment = datetime.fromtimestamp(int(m.group(1)) / 1000, ZoneInfo(tz))
    elif m := DATE_STAMP.search(name):
        moment = datetime.strptime(m.group(1), "%Y-%m-%d").replace(tzinfo=ZoneInfo(tz))
    else:
        return None
    return moment.utcoffset()


def _files(path: Path) -> list[Path]:
    if path.is_dir():
        return sorted(p for p in path.rglob("*") if p.suffix.lower() in (".xlsx", ".xlsm"))
    return [path]


def _load(path: Path, tz: str):
    tables = [(f, *_table(f)) for f in _files(path)]
    labelled = {}
    for f, head, _ in tables:
        if (off := _label_offset(head)) is not None and (m := RANGE.search(f.name)):
            labelled[m.group(1)] = off
    out, skipped = [], {}
    for f, head, rows in tables:
        cols, eng = _english(head, rows)
        role = _role(cols)
        if role is None:
            if rows:
                skipped[f.name] = len(rows)
            continue
        off = _label_offset(head)
        if off is None:
            off = _name_offset(f.name, tz)
        if off is None and (m := RANGE.search(f.name)):
            off = labelled.get(m.group(1))
        if off is None:
            raise ValueError(
                f"{f.name}: no time zone - MEXC writes the time of the download "
                "moment; neither the columns nor the file name tell it"
            )
        out.append((role, timezone(off), eng))
    return out, skipped


def skipped(path: Path, options: dict | None = None) -> dict[str, int]:
    return _load(path, (options or {}).get("timezone", "Europe/Berlin"))[1]


def read(path: Path, options: dict | None = None) -> Iterator[Movement]:
    tables, _ = _load(path, (options or {}).get("timezone", "Europe/Berlin"))
    handlers = {
        "trades": _trades,
        "deposits": _deposits,
        "withdrawals": _withdrawals,
        "otc": _otc,
        "transfers": _transfers,
        "futures_flow": _futures_flow,
        "positions": _positions,
    }
    for role, tz, rows in tables:
        yield from handlers[role](rows, tz)


def _num(raw: str | None) -> Decimal | None:
    v = (raw or "").strip()
    if not v or v == "-":
        return None
    try:
        return Decimal(v)
    except InvalidOperation as exc:
        raise ValueError(f"not a number: {raw!r}") from exc


def _amount_unit(raw: str | None) -> tuple[Decimal | None, str | None]:
    v = (raw or "").strip()
    if not v or v == "-":
        return None, None
    m = AMOUNT_UNIT.match(v)
    if not m:
        raise ValueError(f"unreadable amount {raw!r}")
    n = _num(m.group(1))
    return (abs(n), m.group(2) or None) if n else (None, None)


def _time(raw: str, tz: timezone) -> datetime:
    text = (raw or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=tz)
        except ValueError:
            continue
    raise ValueError(f"unreadable timestamp {raw!r}")


def _clean(r: dict) -> dict:
    return {k: v for k, v in r.items() if v and k != "UID"}


def _ok(r: dict) -> bool:
    status = (r.get("Status") or "").lower()
    return not status or any(s in status for s in SUCCESS)


def _trades(rows, tz) -> Iterator[Movement]:
    for r in rows:
        base, _, quote = (r.get("Pairs") or "").strip().upper().partition("_")
        if not quote:
            raise ValueError(f"unreadable pair {r.get('Pairs')!r}")
        qty, total = _num(r.get("Executed Amount")), _num(r.get("Total"))
        fee, fee_asset = _amount_unit(r.get("Fee"))
        buy = (r.get("Side") or "").strip().lower() in ("buy", "kaufen")
        yield Movement(
            timestamp=_time(r.get("Time", ""), tz),
            kind=Kind.TRADE,
            incoming_asset=base if buy else quote,
            incoming_amount=qty if buy else total,
            outgoing_asset=quote if buy else base,
            outgoing_amount=total if buy else qty,
            fee_asset=(fee_asset or quote) if fee else None,
            fee_amount=fee,
            comment=f"{r.get('Pairs')} {r.get('Side')}",
            raw=_clean(r),
        )


def _deposits(rows, tz) -> Iterator[Movement]:
    for r in rows:
        if not _ok(r):
            continue
        yield Movement(
            timestamp=_time(r.get("Time", ""), tz),
            kind=Kind.RECEIVE,
            incoming_asset=r.get("Crypto"),
            incoming_amount=_num(r.get("Deposit Amount")),
            external_id=(r.get("TxID") or "").split(":")[0].strip() or None,
            comment=r.get("Network") or None,
            raw=_clean(r),
        )


def _withdrawals(rows, tz) -> Iterator[Movement]:
    for r in rows:
        if not _ok(r):
            continue
        coin = (r.get("Crypto") or "").strip()
        fee, settled, requested = (
            _num(r.get("Trading Fee")),
            _num(r.get("Settlement Amount")),
            _num(r.get("Request Amount")),
        )
        if None not in (fee, settled, requested) and settled + fee != requested:
            raise ValueError(
                f"withdrawal does not add up: {requested} - {fee} != {settled} ({coin})"
            )
        yield Movement(
            timestamp=_time(r.get("Time", ""), tz),
            kind=Kind.SEND,
            outgoing_asset=coin,
            outgoing_amount=settled if settled is not None else requested,
            fee_asset=coin if fee else None,
            fee_amount=fee or None,
            external_id=(r.get("TxID") or "").strip() or None,
            counterparty=(r.get("Withdrawal Address") or "").strip() or None,
            comment=r.get("Network") or None,
            raw=_clean(r),
        )


def _otc(rows, tz) -> Iterator[Movement]:
    for r in rows:
        if (r.get("Status") or "").strip().lower() != "complete":
            continue
        buy = (r.get("Trading Direction") or "").strip().lower() == "buy"
        token, fiat = (
            (r.get("Trading Token") or "").strip(),
            (r.get("Settlement Token") or "").strip(),
        )
        qty, amount = _num(r.get("Order Quantity")), _num(r.get("Order Amount"))
        yield Movement(
            timestamp=_time(r.get("End Time") or r.get("Start Time", ""), tz),
            kind=Kind.TRADE,
            incoming_asset=token if buy else fiat,
            incoming_amount=qty if buy else amount,
            outgoing_asset=fiat if buy else token,
            outgoing_amount=amount if buy else qty,
            external_id=(r.get("Order ID") or "").strip() or None,
            comment=r.get("Payment Method") or None,
            raw=_clean(r),
        )


def _transfers(rows, tz) -> Iterator[Movement]:
    for r in rows:
        if not _ok(r):
            continue
        src, dst = (r.get("From System") or "").upper(), (r.get("To System") or "").upper()
        coin, qty = (r.get("Currency") or "").strip(), _num(r.get("Amount"))
        spotlike = {"SPOT", "FIAT"}
        if src in spotlike and dst in spotlike:
            sides = {
                "incoming_asset": coin,
                "incoming_amount": qty,
                "outgoing_asset": coin,
                "outgoing_amount": qty,
            }  # within spot: net zero
        elif src in spotlike:
            sides = {"outgoing_asset": coin, "outgoing_amount": qty}
        elif dst in spotlike:
            sides = {"incoming_asset": coin, "incoming_amount": qty}
        else:
            continue  # between non-spot accounts; their side is in their own ledger
        yield Movement(
            timestamp=_time(r.get("create_time") or r.get("update_time", ""), tz),
            kind=Kind.INTERNAL_TRANSFER,
            comment=f"{src} -> {dst}",
            raw=_clean(r),
            **sides,
        )


FLOW_KINDS = {
    "FEE": Kind.FEE,
    "FUNDING": Kind.FUNDING,
    "CLOSE_POSITION": Kind.REALIZED_PNL,
    "LIQUIDATION": Kind.REALIZED_PNL,
    "TRANSFER": Kind.INTERNAL_TRANSFER,
}


def _futures_flow(rows, tz) -> Iterator[Movement]:
    for r in rows:
        typ = (r.get("Fund Type") or "").strip().upper()
        if typ == "BONUS_DEDUCT":
            continue  # the bonus paying a loss that is booked anyway (see module doc)
        coin, amount = (r.get("Crypto") or "").strip(), _num(r.get("Amount")) or Decimal(0)
        kind = FLOW_KINDS.get(typ, Kind.OTHER)
        extra = {}
        if typ == "BONUS" and amount > 0:
            kind, extra = Kind.INCOME, {"income_kind": IncomeKind.REWARD}
        elif typ == "BONUS":
            # an unused bonus running out: a futures cost that undoes its income (it was
            # left as an open send before, 2026-10-11)
            kind = Kind.FEE
        if kind is Kind.FEE:
            sides = {"fee_asset": coin, "fee_amount": abs(amount)}
        elif amount >= 0:
            sides = {"incoming_asset": coin, "incoming_amount": amount}
        else:
            sides = {"outgoing_asset": coin, "outgoing_amount": -amount}
        # the German export writes POL_USDT, the English one POLUSDT: one spelling
        instrument = (r.get("Futures Trading Pair") or "").strip().replace("_", "") or None
        yield Movement(
            timestamp=_time(r.get("Time", ""), tz),
            kind=kind,
            market=Market.FUTURES,
            instrument=instrument,
            comment=typ,
            raw=_clean(r),
            **sides,
            **extra,
        )


MARGIN = {"Isoliert": "Isolated", "Cross": "Cross"}


def _positions(rows, tz) -> Iterator[Movement]:
    for r in rows:
        close = r.get("Close Time") or r.get("Open Time", "")
        mode = r.get("Margin Mode")
        yield Movement(
            timestamp=_time(close, tz),
            kind=Kind.FUTURES_POSITION,
            market=Market.FUTURES,
            instrument=(r.get("Futures") or "").strip(),
            comment=f"{r.get('Direction')} {MARGIN.get(mode, mode)}".strip(),
            raw=_clean(r),
        )
