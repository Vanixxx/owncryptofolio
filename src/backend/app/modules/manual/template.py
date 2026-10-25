"""The xlsx template for many manual entries at once (ADR 0036).

Columns are found by their heading (German or English), not by position: whoever
inserts a column must not break the file. Times are local time in the configured
zone, like in the form. A row that cannot be read is named with its row number and
reason, never guessed.
"""

from __future__ import annotations

import io
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.core.errors import AppError
from app.modules.manual.service import EntryIn, to_movement
from app.modules.plugins.contract import Movement

COLUMNS = {  # key: (German, English)
    "date": ("Datum", "Date"),
    "time": ("Uhrzeit", "Time"),
    "market": ("Markt", "Market"),
    "kind": ("Art", "Type"),
    "income_kind": ("Ertragsart", "Income type"),
    "incoming_amount": ("Erhalten Menge", "Received amount"),
    "incoming_asset": ("Erhalten Asset", "Received asset"),
    "outgoing_amount": ("Abgegeben Menge", "Spent amount"),
    "outgoing_asset": ("Abgegeben Asset", "Spent asset"),
    "fee_amount": ("Gebühr Menge", "Fee amount"),
    "fee_asset": ("Gebühr Asset", "Fee asset"),
    "instrument": ("Kontrakt", "Contract"),
    "comment": ("Notiz", "Note"),
}
KINDS = {  # kind: (German, English) — the same words as the form
    "trade": ("Kauf/Verkauf", "Buy/Sell"),
    "receive": ("Krypto-Einzahlung", "Crypto deposit"),
    "send": ("Krypto-Auszahlung", "Crypto withdrawal"),
    "fiat_deposit": ("Fiat-Einzahlung", "Fiat deposit"),
    "fiat_withdrawal": ("Fiat-Auszahlung", "Fiat withdrawal"),
    "income": ("Ertrag", "Income"),
    "fee": ("Gebühr", "Fee"),
    "internal_transfer": ("Umbuchung", "Transfer"),
    "migration": ("Umstellung", "Token swap 1:1"),
    "realized_pnl": ("Realisiertes Ergebnis", "Realized PnL"),
    "funding": ("Funding", "Funding"),
}
MARKETS = {"spot": ("Spot", "Spot"), "futures": ("Futures", "Futures")}
INCOME_KINDS = {
    "staking": ("Staking", "Staking"),
    "reward": ("Belohnung", "Reward"),
    "interest": ("Zinsen", "Interest"),
    "airdrop": ("Airdrop", "Airdrop"),
}
EXAMPLE = {
    "date": date(2024, 3, 15),
    "time": time(14, 5, 30),
    "market": "spot",
    "kind": "trade",
    "incoming_amount": "0.5",
    "incoming_asset": "ETH",
    "outgoing_amount": "1500",
    "outgoing_asset": "USDT",
    "fee_amount": "1.5",
    "fee_asset": "USDT",
}
HELP = {
    "de": [
        "So füllst du die Vorlage aus",
        "Eine Zeile = ein Vorgang. Die Beispielzeile vor dem Hochladen löschen.",
        "Datum und Uhrzeit in deiner Ortszeit ({tz}).",
        "Mengen immer positiv. Die Richtung ergibt sich aus Erhalten/Abgegeben.",
        "Markt: Spot (leer = Spot) oder Futures. Bei Futures: Realisiertes Ergebnis, Funding, "
        "Gebühr oder Umbuchung; Gewinn bzw. hinein als Erhalten, Verlust bzw. hinaus als "
        "Abgegeben; Kontrakt optional (z. B. BTCUSDT).",
        "Art: " + ", ".join(v[0] for v in KINDS.values()),
        "Ertragsart (nur bei Ertrag): " + ", ".join(v[0] for v in INCOME_KINDS.values()),
        "Kauf/Verkauf braucht Erhalten und Abgegeben, "
        "Einzahlung nur Erhalten, Auszahlung nur Abgegeben.",
        "Dieselbe Datei zweimal hochladen legt nichts doppelt an.",
    ],
    "en": [
        "How to fill in the template",
        "One row = one transaction. Delete the example row before uploading.",
        "Date and time in your local time ({tz}).",
        "Amounts are always positive. Direction comes from received/spent.",
        "Market: Spot (empty = Spot) or Futures. Futures: Realized PnL, Funding, Fee or "
        "Transfer; a gain or transfer in as received, a loss or transfer out as spent; "
        "contract optional (e.g. BTCUSDT).",
        "Type: " + ", ".join(v[1] for v in KINDS.values()),
        "Income type (income only): " + ", ".join(v[1] for v in INCOME_KINDS.values()),
        "Buy/Sell needs received and spent, a deposit only received, a withdrawal only spent.",
        "Uploading the same file twice adds nothing twice.",
    ],
}


def write(lang: str, tz: str) -> bytes:
    i = 0 if lang == "de" else 1
    wb = Workbook()
    ws = wb.active
    ws.title = "Vorgaenge" if i == 0 else "Transactions"
    ws.append([v[i] for v in COLUMNS.values()])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    row = []
    for key in COLUMNS:
        value = EXAMPLE.get(key)
        if key == "kind" and value:
            value = KINDS[value][i]
        elif key == "market" and value:
            value = MARKETS[value][i]
        row.append(value)
    ws.append(row)
    ws.cell(2, 1).number_format = "yyyy-mm-dd"
    ws.cell(2, 2).number_format = "hh:mm:ss"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)
    help_sheet = wb.create_sheet("Hilfe" if i == 0 else "Help")
    for line in HELP["de" if i == 0 else "en"]:
        help_sheet.append([line.format(tz=tz)])
    help_sheet["A1"].font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 100
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _norm(text) -> str:
    return " ".join(str(text or "").replace("-", " ").replace("/", " ").split()).casefold()


def _lookup(table: dict[str, tuple[str, str]]) -> dict[str, str]:
    out = {}
    for key, (de, en) in table.items():
        for name in (key, de, en):
            out[_norm(name)] = key
    return out


KIND_NAMES, INCOME_NAMES, MARKET_NAMES = _lookup(KINDS), _lookup(INCOME_KINDS), _lookup(MARKETS)
HEADINGS = {_norm(n): key for key, names in COLUMNS.items() for n in names}


def _time_text(d, t) -> str:
    if isinstance(d, datetime):
        day, clock = d.date(), d.time() if t in (None, "") else None
    elif isinstance(d, date):
        day, clock = d, None
    else:
        text = str(d or "").strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
            try:
                day = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError("date")
        clock = None
    if clock is None:
        if isinstance(t, time):
            clock = t
        elif isinstance(t, datetime):
            clock = t.time()
        elif t in (None, ""):
            clock = time(0, 0)
        else:
            text = str(t).strip()
            for fmt in ("%H:%M:%S", "%H:%M"):
                try:
                    clock = datetime.strptime(text, fmt).time()
                    break
                except ValueError:
                    continue
            else:
                raise ValueError("time")
    return datetime.combine(day, clock.replace(microsecond=0)).isoformat()


def read(path: Path, tz: str) -> tuple[list[Movement], list[dict], int]:
    """Movements, problems ({row, code, params} — translated in the UI), rejected rows."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — not an xlsx at all
        raise AppError("manual.template_unreadable", 400) from exc
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    heading = next(rows, None) or ()
    cols = {HEADINGS[_norm(h)]: i for i, h in enumerate(heading) if _norm(h) in HEADINGS}
    if missing := {"date", "kind"} - cols.keys():
        raise AppError("manual.template_columns", 400, column=COLUMNS[sorted(missing)[0]][0])
    movements: list[Movement] = []
    problems: list[dict] = []
    seen: Counter[tuple] = Counter()
    rejected = 0
    for number, values in enumerate(rows, start=2):
        cell = {k: (values[i] if i < len(values) else None) for k, i in cols.items()}
        if all(v in (None, "") for v in cell.values()):
            continue
        try:
            kind = KIND_NAMES.get(_norm(cell.get("kind")))
            if kind is None:
                raise ValueError("kind")
            income = cell.get("income_kind")
            market = MARKET_NAMES.get(_norm(cell.get("market")) or "spot")
            if market is None:
                raise ValueError("market")
            instrument = cell.get("instrument")
            entry = EntryIn(
                kind=kind,
                market=market,
                instrument=str(instrument) if instrument not in (None, "") else None,
                income_kind=INCOME_NAMES.get(_norm(income)) if income else None,
                time=_time_text(cell.get("date"), cell.get("time")),
                comment=str(cell["comment"]) if cell.get("comment") not in (None, "") else None,
                **{
                    k: (None if cell.get(k) in (None, "") else _cell_text(cell[k]))
                    for k in COLUMNS
                    if k.endswith(("_asset", "_amount"))
                },
            )
            content = tuple(sorted(entry.model_dump().items()))
            seen[content] += 1
            # n-th equal row in this file: two real equal buys stay two, a second
            # upload of the same file is recognised (ADR 0036)
            movements.append(to_movement(entry, tz, f"template:{seen[content]}"))
        except AppError as exc:
            rejected += 1
            problems.append({"row": number, "code": exc.code, "params": exc.params})
        except ValueError as exc:
            rejected += 1
            known = ("kind", "date", "time", "market")
            what = exc.args[0] if exc.args and exc.args[0] in known else "row"
            problems.append({"row": number, "code": f"manual.invalid_{what}", "params": {}})
    return movements, problems, rejected


def _cell_text(v) -> str:
    """A number cell comes as a float: Excel keeps 15 significant digits, so its noise
    (0.1 + 0.2 = 0.30000000000000004) goes away there (review 2026-10-24)."""
    if isinstance(v, float):
        return format(Decimal(f"{v:.15g}").normalize(), "f")
    return str(v)
