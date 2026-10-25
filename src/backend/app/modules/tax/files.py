"""The tax report as files (ADR 0042): Excel and PDF, both with the summary and every line.

German only, like the report it serves: it goes to a German tax adviser. Every file names
its method and a plain note for the reader.
"""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.sheets import as_text
from app.modules.tax.lots import BERLIN

METHOD = {"fifo": "FIFO (älteste zuerst)", "lifo": "LIFO (jüngste zuerst)"}
HOW = {
    "trade": "Kauf/Verkauf/Tausch",
    "external": "an Fremde (Zahlung)",
    "fee": "Gebühr in Coin",
    "futures": "Futures-Konto",
}
UNCOVERED = {
    "no_lot": "Verkauf ohne erfassten Kauf",
    "transfer": "Transfer ohne Bestand",
    "migration": "Umstellung ohne Bestand",
    "open_send": "Abgang ohne Gegenstück",
}
FUTURES = {"pnl": "Ergebnis", "funding": "Funding", "fee": "Gebühr", "bonus": "Bonus der Börse"}
INCOME = {
    "staking": "Staking",
    "reward": "Belohnungen",
    "interest": "Zinsen",
    "airdrop": "Airdrops",
}
FORM = {
    "s23": "§ 23 private Veräußerungsgeschäfte (Saldo)",
    "s22": "§ 22 Nr. 3 sonstige Einkünfte (steuerrelevant)",
    "kap19": "ausländische Kapitalerträge",
    "kap21": "darin Gewinne aus Termingeschäften",
    "kap24": "Verluste aus Termingeschäften",
}
OUTSIDE = "außerhalb (eigene Wallet, nicht erfasst)"
# for whoever reads the file (tax adviser, tax office): no internal references (Vanixxx, 2026-10-23)
NOTE = (
    "Erstellt mit owncryptofolio aus den erfassten Transaktionen. Die Aufstellung ist eine "
    "Auswertung und keine Steuerberatung; die Angaben sind vor der Verwendung zu prüfen."
)


def _d(v: str | None) -> Decimal | None:
    return None if v is None else Decimal(v)


def _t(v: str | None) -> datetime | None:
    return None if v is None else datetime.fromisoformat(v).astimezone(BERLIN).replace(tzinfo=None)


def _yes(v: bool) -> str:
    return "ja" if v else "nein"


def holder_lines(rep: dict) -> list[tuple[str, str]]:
    """Name, address and tax number from the settings, only those filled in (Vanixxx,
    2026-10-24); kept reports from before have none."""
    h = rep["summary"].get("holder") or {}
    address = ", ".join(x.strip() for x in (h.get("address") or "").splitlines() if x.strip())
    rows = [("Name", h.get("name") or ""), ("Anschrift", address)]
    rows.append(("Steuernummer", h.get("tax_number") or ""))
    return [(label, value) for label, value in rows if value]


def summary_lines(rep: dict) -> list[tuple[str, object]]:
    s = rep["summary"]
    s23, s22, s20 = s["s23"], s["s22"], s["s20"]
    out: list[tuple[str, object]] = [
        ("Steuerjahr", s["year"]),
        ("Verfahren", METHOD[s["method"]]),
        ("Betrachtung", "walletbezogen: je Wallet/Börse, Spot und Futures getrennt"),
        ("Airdrops versteuert", _yes(s["airdrop_taxable"])),
        ("", None),
    ]
    if s.get("forms"):
        out.append(("Ausfüllhilfe (Zeilen laut Formular des Jahres prüfen)", None))
        for f in s["forms"]:
            where = f"Anlage {f['form']}" + (f" Zeile {f['line']}" if f["line"] else "")
            out.append((f"{where}: {FORM[f['key']]}", _d(f["value"])))
        out.append(("", None))
    out += [
        ("§ 23 EStG – private Veräußerungsgeschäfte", None),
        ("Gewinne (innerhalb der Haltefrist)", _d(s23["gains"])),
        ("Verluste (innerhalb der Haltefrist)", _d(s23["losses"])),
        ("Saldo", _d(s23["net"])),
        ("Freigrenze", _d(s23["limit"])),
        ("Saldo steuerpflichtig (Freigrenze erreicht)", _yes(s23["taxable"])),
        ("steuerfreie Gewinne (Haltefrist über 1 Jahr)", _d(s23["tax_free_gain"])),
        ("Posten / davon in der Frist", f"{s23['count']} / {s23['count_taxed']}"),
        ("Posten mit geschätzter Kostenbasis", s23["estimated"]),
        ("Posten ohne Kurs", s23["unvalued"]),
        ("", None),
        ("§ 22 Nr. 3 EStG – Erträge (Staking, Belohnungen)", None),
    ]
    out += [(f"davon {INCOME.get(k, k)}", _d(v)) for k, v in s22.get("by_kind", {}).items()]
    out += [
        ("Summe", _d(s22["total"])),
        ("Freigrenze (steuerfrei unter)", _d(s22["limit"])),
        ("steuerpflichtig", _yes(s22["taxable"])),
        ("", None),
        ("§ 20 EStG – Futures (Termingeschäfte, Anlage KAP)", None),
        ("Gewinne", _d(s20["gains"])),
        ("Verluste", _d(s20["losses"])),
        ("Funding erhalten", _d(s20["funding_in"])),
        ("Funding gezahlt", _d(s20["funding_out"])),
        ("Gebühren", _d(s20["fees"])),
    ]
    if "bonus" in s20:
        out.append(("Bonus der Börse (gutgeschrieben minus verfallen)", _d(s20["bonus"])))
    out.append(("Saldo", _d(s20["net"])))
    if "kap_gains" in s20:
        out += [
            ("Gewinne für Anlage KAP (Ergebnisse, Funding, Bonus)", _d(s20["kap_gains"])),
            ("Verluste für Anlage KAP (Ergebnisse, Funding, Gebühren)", _d(s20["kap_losses"])),
        ]
    c = (s.get("carry") or {}).get("year")
    if c:
        out += [
            ("", None),
            ("Verlustvortrag § 23 (Vorschlag, maßgeblich ist der Bescheid)", None),
            ("Vortrag zu Jahresbeginn", _d(c["start"])),
            ("mit dem Gewinn des Jahres verrechnet", _d(c["used"])),
            ("Rücktrag aus dem Folgejahr", _d(c["back_in"])),
            ("Verlust des Jahres, ins Vorjahr zurückgetragen", _d(c["back_out"])),
            ("Verlust des Jahres, vorgetragen", _d(c["loss"])),
            ("Gewinn nach Verrechnung", _d(c["remaining"])),
            ("Vortrag zum Jahresende (gerechnet)", _d(c["end"])),
        ]
        if c["assessed"] is not None:
            out.append(("Vortrag zum Jahresende laut Bescheid", _d(c["assessed"])))
    out += [
        ("", None),
        ("Verluste durch Pleite/Betrug (nicht verrechnet)", _d(s["lost"]["cost"])),
        (
            "Offene Punkte (ungedeckt, ohne Gegenstück)",
            s["open"]["uncovered"] + s["open"]["open_sends"],
        ),
        ("", None),
        ("Hinweis", NOTE),
    ]
    return out


ESTIMATED = (
    "„Kosten geschätzt“ = ja: Für diese Coins steht kein Kauf in den Daten (z. B. ein "
    "Eingang ohne Gegenstück). Als Kosten gilt ersatzweise ihr Marktwert beim Eingang."
)

# column kinds: text · dt (date and time) · day · int · qty · eur
Column = tuple[str, str, float]  # header, kind, width in the PDF (mm, landscape)


def tables(rep: dict) -> list[tuple[str, list[Column], list[list]]]:
    """The detail sheets, once for Excel and PDF alike."""
    lines, s = rep["lines"], rep["summary"]
    place = lambda p: p or OUTSIDE  # noqa: E731
    return [
        (
            "Summen je Quelle",
            [
                ("Ort", "text", 70),
                ("Gewinne €", "eur", 25),
                ("Verluste €", "eur", 25),
                ("Saldo €", "eur", 25),
                ("steuerfrei Gewinne €", "eur", 30),
                ("steuerfrei Verluste €", "eur", 30),
                ("steuerfrei Saldo €", "eur", 30),
            ],
            [
                [
                    place(b["place"]),
                    _d(b["gains"]),
                    _d(b["losses"]),
                    _d(b["net"]),
                    _d(b["free_gains"]),
                    _d(b["free_losses"]),
                    _d(b["free_net"]),
                ]
                for b in lines.get("by_source", [])
            ],
        ),
        (
            "§ 23 Veräußerungen",
            [
                ("Coin", "text", 16),
                ("Ort", "text", 52),
                ("Verkauft", "dt", 26),
                ("Angeschafft", "dt", 26),
                ("Tage", "int", 11),
                ("Menge", "qty", 26),
                ("Erlös €", "eur", 20),
                ("Kosten €", "eur", 20),
                ("Gewinn €", "eur", 20),
                ("steuerfrei", "text", 14),
                ("Kosten geschätzt", "text", 22),
                ("Vorgang", "text", 22),
                ("Zeile", "int", 12),
                ("Kauf-Zeile", "int", 14),
            ],
            [
                [
                    d["asset"],
                    place(d["place"]),
                    _t(d["sold"]),
                    _t(d["acquired"]),
                    d["days"],
                    _d(d["qty"]),
                    _d(d["proceeds"]),
                    _d(d["cost"]),
                    _d(d["gain"]),
                    _yes(d["tax_free"]),
                    _yes(d["estimated"]),
                    HOW.get(d["how"], d["how"]),
                    d["row"],
                    d.get("origin") or None,
                ]
                for d in lines["disposals"]
            ],
        ),
        (
            "§ 22 Erträge",
            [
                ("Coin", "text", 20),
                ("Ort", "text", 70),
                ("Zufluss", "dt", 30),
                ("Menge", "qty", 35),
                ("Wert €", "eur", 25),
                ("Art", "text", 30),
                ("Zeile", "int", 15),
            ],
            [
                [
                    i["asset"],
                    place(i["place"]),
                    _t(i["at"]),
                    _d(i["qty"]),
                    _d(i["value"]),
                    i["kind"],
                    i["row"],
                ]
                for i in lines["incomes"]
            ],
        ),
        (
            "§ 20 Futures",
            [
                ("Ort", "text", 70),
                ("Zeit", "dt", 30),
                ("Art", "text", 25),
                ("Coin", "text", 20),
                ("Menge", "qty", 35),
                ("Wert €", "eur", 25),
                ("Zeile", "int", 15),
            ],
            [
                [
                    place(f["place"]),
                    _t(f["at"]),
                    FUTURES.get(f["kind"], f["kind"]),
                    f["asset"],
                    _d(f["amount"]),
                    _d(f["value"]),
                    f["row"],
                ]
                for f in lines["futures"]
            ],
        ),
        (
            "Verluste Pleite-Betrug",
            [
                ("Coin", "text", 20),
                ("Ort", "text", 70),
                ("Verloren", "dt", 30),
                ("Angeschafft", "dt", 30),
                ("Menge", "qty", 35),
                ("Kosten €", "eur", 25),
                ("Zeile", "int", 15),
            ],
            [
                [
                    x["asset"],
                    place(x["place"]),
                    _t(x["at"]),
                    _t(x["acquired"]),
                    _d(x["qty"]),
                    _d(x["cost"]),
                    x["row"],
                ]
                for x in lines["losses"]
            ],
        ),
        (
            "Offene Punkte",
            [
                ("Coin", "text", 20),
                ("Ort", "text", 70),
                ("Zeit", "dt", 30),
                ("Menge", "qty", 35),
                ("Erlösanteil €", "eur", 25),
                ("Grund", "text", 50),
                ("Zeile", "int", 15),
            ],
            [
                [
                    u["asset"],
                    place(u["place"]),
                    _t(u["at"]),
                    _d(u["qty"]),
                    _d(u["proceeds"]),
                    UNCOVERED.get(u["reason"], u["reason"]),
                    u["row"],
                ]
                for u in lines["uncovered"]
            ],
        ),
        (
            "Bestandsveränderungen",
            [
                ("Ort", "text", 60),
                ("Coin", "text", 18),
                ("Bestand 01.01.", "qty", 34),
                ("Zufluss", "qty", 34),
                ("Abfluss", "qty", 34),
                ("Bestand 31.12.", "qty", 34),
                ("Wert 31.12. €", "eur", 25),
            ],
            [
                [
                    place(h["place"]),
                    h["asset"],
                    _d(h["start"]),
                    _d(h["in"]),
                    _d(h["out"]),
                    _d(h["end"]),
                    _d(h["value"]),
                ]
                for h in lines.get("holdings", [])
            ],
        ),
        (
            "Bestand 31.12.",
            [
                ("Coin", "text", 20),
                ("Ort", "text", 70),
                ("Markt", "text", 18),
                ("Menge", "qty", 35),
                ("Angeschafft", "dt", 30),
                ("Kosten €", "eur", 25),
                ("steuerfrei nach", "day", 25),
                ("Kosten geschätzt", "text", 22),
            ],
            [
                [
                    b["asset"],
                    place(b["place"]),
                    b["market"],
                    _d(b["qty"]),
                    _t(b["acquired"]),
                    _d(b["cost"]),
                    datetime.fromisoformat(b["free_from"]),
                    _yes(b["estimated"]),
                ]
                for b in lines["year_end"]
            ],
        ),
        (
            "Verlustvortrag",
            [
                ("Jahr", "int", 12),
                ("Saldo § 23 €", "eur", 25),
                ("Vortrag Anfang €", "eur", 26),
                ("verrechnet €", "eur", 22),
                ("Rücktrag erhalten €", "eur", 27),
                ("Rücktrag ins Vorjahr €", "eur", 30),
                ("Verlust vorgetragen €", "eur", 29),
                ("Vortrag Ende €", "eur", 24),
                ("laut Bescheid €", "eur", 24),
                ("Gewinn nach Verrechnung €", "eur", 34),
            ],
            [
                [
                    r["year"],
                    _d(r["net"]),
                    _d(r["start"]),
                    _d(r["used"]),
                    _d(r["back_in"]),
                    _d(r["back_out"]),
                    _d(r["loss"]),
                    _d(r["end"]),
                    _d(r["assessed"]),
                    _d(r["remaining"]),
                ]
                for r in (s.get("carry") or {}).get("rows", [])
            ],
        ),
    ]


def xlsx(rep: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Zusammenfassung"
    head = holder_lines(rep)
    for label, value in head:
        ws.append([label, value])
    if head:
        ws.append([])
    for label, value in summary_lines(rep):
        ws.append([label, value])
        if value is None and label:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["Kosten geschätzt", ESTIMATED])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 28

    for title, cols, rows in tables(rep):
        w = wb.create_sheet(title)
        w.append([c[0] for c in cols])
        for c in w[1]:
            c.font = Font(bold=True)
        for r in rows:
            w.append(r)
        for i, (_, kind, _) in enumerate(cols, 1):
            w.column_dimensions[get_column_letter(i)].width = 18
            if kind in ("dt", "day"):
                for (cell,) in w.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.number_format = "dd.mm.yyyy hh:mm" if kind == "dt" else "dd.mm.yyyy"
        w.freeze_panes = "A2"
    for sheet in wb.worksheets:
        as_text(sheet)  # outside text never becomes a formula (review 2026-10-24)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v: Decimal, places: int | None) -> str:
    text = f"{v:,.{places}f}" if places is not None else f"{v.normalize():,f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def _cell(v: object, kind: str) -> str:
    if v is None:
        return ""
    if kind == "dt" and isinstance(v, datetime):
        return f"{v:%d.%m.%Y %H:%M}"
    if kind == "day" and isinstance(v, datetime):
        return f"{v:%d.%m.%Y}"
    if kind == "eur" and isinstance(v, Decimal):
        return _num(v, 2)
    if kind == "qty" and isinstance(v, Decimal):
        return _num(v.quantize(Decimal("0.00000001")) if v.as_tuple().exponent < -8 else v, None)
    return str(v)


def _de(v: object) -> str:
    """German number format; EUR as text (the built-in PDF fonts have no euro sign)."""
    if isinstance(v, Decimal):
        text = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"{text} EUR"
    return "" if v is None else str(v)


class _Doc(FPDF):
    def normalize_text(self, text: str) -> str:
        # the built-in fonts know Latin-1 only: dashes become hyphens, anything else "?"
        text = (
            text.replace("–", "-")
            .replace("×", "x")
            .replace("€", "EUR")
            .replace("„", '"')
            .replace("“", '"')
        )
        return super().normalize_text(text.encode("latin-1", "replace").decode("latin-1"))

    def footer(self) -> None:
        self.set_y(-10)
        self.set_font("Helvetica", "", 7)
        self.cell(0, 5, f"owncryptofolio · {self.title_text} · Seite {self.page_no()}", align="C")


def _fit(doc: FPDF, text: str, width: float) -> str:
    """Cuts what does not fit its column, with an ellipsis (long wallet names)."""
    if doc.get_string_width(text) <= width - 1:
        return text
    while text and doc.get_string_width(text + "...") > width - 1:
        text = text[:-1]
    return text + "..."


FOOTNOTE = (
    "Kosten geschätzt: Für diese Coins steht kein Kauf in den Daten (z. B. ein Eingang ohne "
    "Gegenstück). Als Kosten gilt ersatzweise ihr Marktwert beim Eingang."
)


def _star_estimated(cols: list[Column], rows: list[list]) -> tuple[list[Column], list[list], bool]:
    """In the PDF the yes/no column gives way to a star at the cost (Vanixxx, 2026-10-13); the
    footnote follows only a table that has one. Excel keeps the column to filter by."""
    names = [c[0] for c in cols]
    if "Kosten geschätzt" not in names:
        return cols, rows, False
    flag, cost = names.index("Kosten geschätzt"), names.index("Kosten €")
    starred = False
    out = []
    for r in rows:
        r = list(r)
        if r[flag] == "ja":
            r[cost] = f"{_cell(r[cost], 'eur')}*"
            starred = True
        out.append(r[:flag] + r[flag + 1 :])
    return cols[:flag] + cols[flag + 1 :], out, starred


def _table(doc: FPDF, cols: list[Column], rows: list[list]) -> None:
    """Plain cells, one line per row, headings again on every page. fpdf2's table() took
    12 s for the 2,500 lines of 2024; this takes well under one."""
    scale = (doc.w - doc.l_margin - doc.r_margin) / sum(w for _, _, w in cols)
    widths = [w * scale for _, _, w in cols]
    right = ["R" if k in ("int", "qty", "eur") else "L" for _, k, _ in cols]
    h = 3.8
    doc.c_margin = 1.2  # a little air between right- and left-aligned neighbours

    def head() -> None:
        doc.set_font("Helvetica", "B", 6.5)
        for (title, _, _), w, a in zip(cols, widths, right, strict=True):
            doc.cell(w, h + 1, _fit(doc, title, w), border="B", align=a)
        doc.ln(h + 1)
        doc.set_font("Helvetica", "", 6.5)

    head()
    for r in rows:
        if doc.get_y() + h > doc.page_break_trigger:
            doc.add_page(orientation="L")
            head()
        for v, (_, kind, _), w, a in zip(r, cols, widths, right, strict=True):
            doc.cell(w, h, _fit(doc, _cell(v, kind), w), align=a)
        doc.ln(h)


def pdf(rep: dict, made: datetime | None = None) -> bytes:
    """Summary first, then every detail sheet of the Excel as a table (Vanixxx, 2026-10-13).
    The summary fits on page one (Vanixxx, 2026-10-24: it used to spill a few lines over)."""
    year = rep["summary"]["year"]
    doc = _Doc(format="A4")
    doc.title_text = f"Steuerauswertung {year} ({rep['summary']['method'].upper()})"
    doc.set_auto_page_break(True, margin=12)
    doc.set_margins(18, 14, 18)
    doc.add_page()
    doc.set_font("Helvetica", "B", 15)
    doc.cell(0, 8, f"Steuerauswertung Krypto {year}", new_x="LMARGIN", new_y="NEXT")
    doc.set_font("Helvetica", "", 9)
    for label, value in holder_lines(rep):
        text = value if label != "Steuernummer" else f"{label}: {value}"
        doc.multi_cell(0, 4.4, text, align="L", new_x="LMARGIN", new_y="NEXT")
    doc.set_font("Helvetica", "", 8.5)
    when = (made or datetime.now(BERLIN)).astimezone(BERLIN)
    doc.cell(0, 5, f"owncryptofolio, erstellt {when:%d.%m.%Y %H:%M}", new_x="LMARGIN", new_y="NEXT")
    doc.ln(2)
    for label, value in summary_lines(rep):
        if not label:
            doc.ln(1.2)
            continue
        if value is None:
            doc.set_font("Helvetica", "B", 10)
            doc.cell(0, 5.0, label, new_x="LMARGIN", new_y="NEXT")
            continue
        doc.set_font("Helvetica", "", 8.5)
        if label == "Hinweis":
            doc.multi_cell(0, 3.9, f"{label}: {value}")
            continue
        doc.cell(130, 3.9, label)
        doc.cell(0, 3.9, _de(value), align="R", new_x="LMARGIN", new_y="NEXT")
    for title, cols, rows in tables(rep):
        cols, rows, starred = _star_estimated(cols, rows)
        doc.add_page(orientation="L")
        doc.set_margins(12, 12, 12)
        doc.set_font("Helvetica", "B", 12)
        doc.cell(0, 8, f"{title} ({len(rows)})", new_x="LMARGIN", new_y="NEXT")
        if not rows:
            doc.set_font("Helvetica", "", 9)
            doc.cell(0, 6, "Keine Einträge in diesem Jahr.", new_x="LMARGIN", new_y="NEXT")
            continue
        _table(doc, cols, rows)
        if starred:
            doc.ln(2)
            doc.set_font("Helvetica", "", 7)
            doc.multi_cell(0, 3.5, f"* {FOOTNOTE}")
    return bytes(doc.output())
