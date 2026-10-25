"""Source of funds as PDF and Excel (ADR 0046): the PDF for the bank, notary or exchange —
summary, split by origin and the chain as a tree; the Excel with every step."""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from app.core.sheets import as_text
from app.modules.tax.files import BERLIN, _cell, _Doc, _fit

KIND = {
    "sale": "Verkauf",
    "swap": "Tausch",
    "payment": "Zahlung an Dritte",
    "transfer": "Transfer",
    "withdrawal": "Auszahlung",
    "buy": "Kauf mit Euro",
    "rebook": "Umbuchung Spot/Futures",
    "holding": "Bestand heute",
    "holdings": "Gesamter Bestand heute",
    "planned": "Geplante Transaktion",
    "bank": "Einzahlung vom Bankkonto",
    "card": "Kauf per Karte",
    "gain": "Kursgewinn",
    "loss": "Kursverlust",
    "income": "Ertrag (Staking, Belohnung, Airdrop)",
    "futures": "Ergebnis aus Futures",
    "unknown": "Zufluss ohne erfassten Kauf",
    "money_unknown": "Euro ohne erfasste Einzahlung",
    "small": "kleine Anteile",
    "more": "weitere Vorgänge",
}
NOTES = [
    "Einzahlung vom Bankkonto / Kauf per Karte: eigenes Geld. Beleg: Kontoauszug bzw. "
    "Kartenabrechnung zum genannten Datum.",
    "Kursgewinn: Wertzuwachs der Coins zwischen Kauf und Verkauf (bzw. Transfer). Beleg: "
    "Kauf- und Verkaufsbelege der Börse (Transaktionsverlauf).",
    "Ertrag: Staking, Belohnungen oder Airdrops der Börse/des Netzwerks. Beleg: Auszug der Börse.",
    "Zufluss ohne erfassten Kauf / Euro ohne erfasste Einzahlung: in den Daten fehlt die "
    "Herkunft, hierfür ist ein eigener Nachweis nötig.",
    "Transfers zwischen eigenen Konten und Wallets sind mit Transaktions-ID aufgeführt; sie "
    "ändern die Herkunft nicht.",
    "Reihenfolge der Bestände: {method}, wie in der Steuerauswertung. Anteile unter 1 % sind im "
    "Baum zusammengefasst; die Excel-Datei enthält jeden Vorgang.",
    "Kursverluste stehen als negativer Anteil da: die Herkunft zeigt das volle eingesetzte Geld, "
    "der Verlust zieht davon ab (die Anteile können zusammen über 100 % ergeben).",
    "Bestand heute: bewertet zum aktuellen Kurs; der Unterschied zu den Kosten ist „Kursgewinn "
    "bzw. Kursverlust (noch nicht verkauft)“.",
    "Datenart: API = von der Börse abgerufen, Blockchain = aus der Kette gelesen, Datei = Export "
    "der Börse eingelesen, von Hand = selbst erfasst (eigener Beleg nötig).",
]
CHANNEL = {"api": "API", "chain": "Blockchain", "file": "Datei", "manual": "von Hand"}
PAID = {
    ("bank", False): "Überweisung",
    ("bank", True): "SEPA direkt",
    ("card", False): "Karte",
    ("card", True): "Karte direkt",
}
SOURCE_KIND = {"exchange": "Börse", "wallet": "Wallet", "outside": "außerhalb"}


def _eur_text(v) -> str:
    return f"{_cell(_d(v), 'eur')} EUR"


def summary(data: dict) -> str:
    """A few sentences from a fixed pattern (Blockpit lets an AI write them; this can be
    checked against the figures below)."""
    t, ov = data["target"], data.get("overview") or {}
    qty = f"{_cell(_d(t['qty']), 'qty')} {t['asset']}" if t.get("qty") else ""
    place, to = t.get("place") or "außerhalb", t.get("to")
    day = _day(t["at"][:10])
    what = {
        "sale": f"wurden {qty} bei {place} verkauft",
        "swap": f"wurden bei {place} {qty} in "
        f"{_cell(_d(t.get('into_qty')), 'qty')} {t.get('into_asset')} getauscht",
        "transfer": f"wurden {qty} von {place} zu {to or 'außerhalb'} übertragen",
        "withdrawal": f"wurden {qty} von {place} ausgezahlt",
        "payment": f"wurden {qty} von {place} an Dritte gezahlt",
        "holding": f"liegen {qty} im Bestand",
        "holdings": "liegt der gesamte Bestand",
        "planned": f"sollen {qty} von {place} abgehen (geplant)",
    }.get(t["kind"], qty)
    out = [f"Am {day} {what}, Wert {_eur_text(t['amount'])}."]
    if t["kind"] == "holdings":
        out = [f"Am {day} hat der gesamte Bestand einen Wert von {_eur_text(t['amount'])}."]
    mix = {m["kind"]: m["value"] for m in data["mix"]}
    own = [
        f"{_eur_text(mix[k])} {label}"
        for k, label in (("bank", "aus Einzahlungen vom Bankkonto"), ("card", "aus Kartenkäufen"))
        if k in mix
    ]
    if own:
        out.append("Eigenes Geld: " + " und ".join(own) + ".")
    more = [
        f"{label} {_eur_text(mix[k])}"
        for k, label in (
            ("gain", "Kursgewinne"),
            ("income", "Erträge"),
            ("futures", "Futures-Ergebnisse"),
            ("loss", "Kursverluste"),
            ("unknown", "ohne erfassten Kauf"),
            ("money_unknown", "Euro ohne erfasste Einzahlung"),
        )
        if k in mix
    ]
    if more:
        out.append("Dazu: " + ", ".join(more) + ".")
    if ov.get("transactions"):
        n = len(ov.get("sources", []))
        out.append(
            f"Die Kette umfasst {ov['transactions']} Transaktionen über {n} "
            f"{'Börse bzw. Wallet' if n == 1 else 'Börsen und Wallets'} "
            f"zwischen {_day(ov['from'][:10])} und {_day(ov['to'][:10])}."
        )
    return " ".join(out)


def _section(doc, title: str) -> None:
    doc.ln(3)
    doc.set_font("Helvetica", "B", 11)
    doc.cell(0, 6.5, title, new_x="LMARGIN", new_y="NEXT")


def _rows(doc, cols: list[tuple[str, float, str]], rows: list[list[str]]) -> None:
    """A small table: header, then one line per row, texts cut to their column."""
    width = doc.w - doc.l_margin - doc.r_margin
    scale = width / sum(w for _, w, _ in cols)
    doc.set_font("Helvetica", "B", 7.5)
    for title, w, a in cols:
        doc.cell(w * scale, 4.6, _fit(doc, title, w * scale), border="B", align=a)
    doc.ln(4.6)
    doc.set_font("Helvetica", "", 7.5)
    for r in rows:
        for v, (_, w, a) in zip(r, cols, strict=True):
            doc.cell(w * scale, 4.3, _fit(doc, v, w * scale), align=a)
        doc.ln(4.3)


def _d(v: str | None) -> Decimal | None:
    return None if v is None else Decimal(v)


def _when(v: str) -> datetime:
    return datetime.fromisoformat(v).astimezone(BERLIN).replace(tzinfo=None)


def _what(n: dict) -> str:
    kind = KIND.get(n["kind"], n["kind"])
    if n["kind"] in ("small", "more"):
        if len(n["mix"]) == 1:  # one kind: its amount stands on the right already
            mix = KIND.get(next(iter(n["mix"])), "")
        else:
            mix = ", ".join(
                f"{KIND.get(k, k)} {_cell(_d(v), 'eur')} EUR" for k, v in n["mix"].items()
            )
        return f"{kind} ({n['count']} Vorgänge): {mix}"
    if n["kind"] == "gain" and n.get("detail") == "until_transfer":
        return "Kursgewinn bis zum Transfer"
    if n["kind"] in ("gain", "loss") and n.get("detail") == "unrealized":
        return f"{KIND[n['kind']]} (noch nicht verkauft)"
    if n["kind"] == "loss" and n.get("detail") == "until_transfer":
        return "Kursverlust bis zum Transfer"
    if n.get("detail") == "direct":
        kind = {"bank": "Bezahlt direkt vom Bankkonto", "card": "Bezahlt direkt per Karte"}[
            n["kind"]
        ]
    elif n.get("detail") and n["kind"] == "income":
        kind = f"Ertrag ({n['detail']})"
    parts = [kind]
    if n.get("qty") and n.get("asset") and n["kind"] not in ("gain",):
        into = (
            f" > {_cell(_d(n['into_qty']), 'qty')} {n['into_asset']}" if n.get("into_asset") else ""
        )
        parts.append(f"{_cell(_d(n['qty']), 'qty')} {n['asset']}{into}")
    where = n.get("place") or ("außerhalb" if "place" in n else None)
    if where:
        parts.append(where + (f" > {n['to']}" if n.get("to") else ""))
    if n.get("at"):
        parts.append(f"{_when(n['at']):%d.%m.%Y}")
    return " · ".join(parts)


def _lines(n: dict, depth: int = 0) -> list[tuple[int, str, Decimal | None, str | None]]:
    out = [(depth, _what(n), _d(n.get("share")), n.get("tx"))]
    for k in n.get("children", []):
        out += _lines(k, depth + 1)
    return out


def _day(v: str) -> str:
    return f"{datetime.fromisoformat(v):%d.%m.%Y}"


def head(data: dict) -> list[tuple[str, str]]:
    t = data["target"]
    out = [
        ("Vorgang", KIND.get(t["kind"], t["kind"])),
        # a planned step names only its day
        ("Datum", f"{_when(t['at']):%d.%m.%Y %H:%M}" if "T" in t["at"] else _day(t["at"])),
    ]
    if t["kind"] not in ("holding", "holdings"):
        out.append(("Ort", (t["place"] or "außerhalb") + (f" > {t['to']}" if t.get("to") else "")))
    if t["kind"] != "holdings":
        into = (
            f" > {_cell(_d(t['into_qty']), 'qty')} {t['into_asset']}" if t.get("into_asset") else ""
        )
        out.append(("Coin / Menge", f"{_cell(_d(t['qty']), 'qty')} {t['asset']}{into}"))
    out.append(
        (
            "Wert" if t["kind"] in ("holding", "holdings") else "Betrag",
            f"{_cell(_d(t['amount']), 'eur')} EUR",
        )
    )
    return out


def _coin_line(c: dict) -> str:
    top = sorted(c["mix"].items(), key=lambda kv: -Decimal(kv[1]))[:3]
    return ", ".join(f"{KIND.get(k, k)} {_cell(_d(v), 'eur')}" for k, v in top)


def pdf(data: dict, holder: tuple[str, str], made: datetime | None = None) -> bytes:
    doc = _Doc(format="A4")
    doc.title_text = "Nachweis der Mittelherkunft"
    doc.set_auto_page_break(True, margin=14)
    doc.set_margins(18, 18, 18)
    doc.add_page()
    doc.set_font("Helvetica", "B", 16)
    doc.cell(0, 10, "Nachweis der Mittelherkunft", new_x="LMARGIN", new_y="NEXT")
    doc.set_font("Helvetica", "", 9)
    when = (made or datetime.now(BERLIN)).astimezone(BERLIN)
    name, address = holder
    if name or address:
        doc.multi_cell(
            0, 4.6, "\n".join(x for x in (name, address) if x), new_x="LMARGIN", new_y="NEXT"
        )
    doc.cell(0, 6, f"owncryptofolio, erstellt {when:%d.%m.%Y %H:%M}", new_x="LMARGIN", new_y="NEXT")
    doc.ln(3)
    doc.set_font("Helvetica", "", 10)
    ov = data.get("overview") or {}
    lines = head(data)
    if ov.get("transactions"):
        lines += [
            ("Zeitraum", f"{_day(ov['from'][:10])} bis {_day(ov['to'][:10])}"),
            ("Transaktionen", str(ov["transactions"])),
            ("Börsen/Wallets", str(len(ov["sources"]))),
            (
                "Datenart",
                ", ".join(f"{CHANNEL.get(k, k)} {v}" for k, v in sorted(ov["channels"].items())),
            ),
        ]
    for label, value in lines:
        doc.cell(40, 5.4, label)
        doc.cell(0, 5.4, value, new_x="LMARGIN", new_y="NEXT")
    _section(doc, "Zusammenfassung")
    doc.set_font("Helvetica", "", 9.5)
    doc.multi_cell(0, 4.8, summary(data), new_x="LMARGIN", new_y="NEXT")
    doc.ln(3)
    doc.set_font("Helvetica", "B", 11)
    doc.cell(0, 6.5, "Herkunft des Betrags", new_x="LMARGIN", new_y="NEXT")
    doc.set_font("Helvetica", "", 10)
    for m in data["mix"]:
        doc.cell(100, 5.4, KIND.get(m["kind"], m["kind"]))
        doc.cell(35, 5.4, f"{_cell(_d(m['value']), 'eur')} EUR", align="R")
        doc.cell(
            0, 5.4, f"{_cell(_d(m['share']), 'eur')} %", align="R", new_x="LMARGIN", new_y="NEXT"
        )
    if data.get("deposits"):
        _section(doc, "Einzahlungen im Detail (eigenes Geld)")
        _rows(
            doc,
            [
                ("Datum", 18, "L"),
                ("Ort", 34, "L"),
                ("Art", 34, "L"),
                ("Betrag", 22, "R"),
                ("davon hier", 22, "R"),
                ("ID", 50, "L"),
                ("Datenart", 18, "L"),
            ],
            [
                [
                    f"{_when(d['at']):%d.%m.%Y}",
                    d["place"] or "außerhalb",
                    PAID[(d["kind"], d.get("detail") == "direct")],
                    _eur_text(d["value"]),
                    _eur_text(d["share"]),
                    d.get("tx") or "",
                    CHANNEL.get(d.get("channel") or "", ""),
                ]
                for d in sorted(data["deposits"], key=lambda d: d["at"])
            ],
        )
    if ov.get("sources"):
        _section(doc, "Datenquellen (beteiligte Börsen und Wallets)")
        _rows(
            doc,
            [
                ("Name", 40, "L"),
                ("Typ", 16, "L"),
                ("Adresse", 80, "L"),
                ("Vorgänge", 16, "R"),
                ("Datenart", 40, "L"),
            ],
            [
                [
                    x["place"] or "außerhalb",
                    SOURCE_KIND.get(x["kind"], x["kind"]),
                    x.get("address") or "",
                    str(x["count"]),
                    ", ".join(f"{CHANNEL.get(k, k)} {v}" for k, v in sorted(x["channels"].items())),
                ]
                for x in ov["sources"]
            ],
        )
    if data.get("coins") and data["target"]["kind"] == "holdings":
        doc.ln(3)
        doc.set_font("Helvetica", "B", 11)
        doc.cell(0, 6.5, "Bestand je Coin", new_x="LMARGIN", new_y="NEXT")
        width = doc.w - doc.l_margin - doc.r_margin
        for c in data["coins"]:
            doc.set_font("Helvetica", "", 8.5)
            doc.cell(18, 4.6, c["asset"])
            doc.cell(28, 4.6, f"{_cell(_d(c['value']), 'eur')} EUR", align="R")
            doc.set_font("Helvetica", "", 7.5)
            doc.cell(
                width - 46,
                4.6,
                _fit(doc, "  " + _coin_line(c), width - 46),
                new_x="LMARGIN",
                new_y="NEXT",
            )
    doc.ln(3)
    doc.set_font("Helvetica", "B", 11)
    doc.cell(0, 6.5, "Belegkette (jüngster Schritt oben)", new_x="LMARGIN", new_y="NEXT")
    width = doc.w - doc.l_margin - doc.r_margin
    for depth, text, share, tx in _lines(data["tree"]):
        indent = min(depth, 12) * 3.2
        doc.set_x(doc.l_margin + indent)
        doc.set_font("Helvetica", "", 8)
        money = "" if share is None else f"{_cell(share, 'eur')} EUR"
        doc.cell(
            width - indent - 32, 4.4, _fit(doc, ("> " if depth else "") + text, width - indent - 32)
        )
        doc.cell(32, 4.4, money, align="R", new_x="LMARGIN", new_y="NEXT")
        if tx:
            doc.set_x(doc.l_margin + indent + 3)
            doc.set_font("Helvetica", "", 6.5)
            doc.cell(
                0, 3.4, _fit(doc, f"Tx {tx}", width - indent - 3), new_x="LMARGIN", new_y="NEXT"
            )
    doc.ln(3)
    doc.set_font("Helvetica", "B", 11)
    doc.cell(0, 6.5, "Erläuterungen und Belege", new_x="LMARGIN", new_y="NEXT")
    doc.set_font("Helvetica", "", 8.5)
    notes = [n.format(method=data["method"].upper()) for n in NOTES]
    if data["target"]["kind"] == "planned":
        notes.insert(
            0,
            "Geplante Transaktion: noch nicht ausgeführt. Die Herkunft beruht auf dem heutigen "
            "Bestand am genannten Ort, in der Reihenfolge der Steuer, bewertet zum heutigen Kurs. "
            "Ändert sich der Bestand bis zur Ausführung, kann sich die Herkunft ändern.",
        )
    if data["lost"]["count"]:
        notes.append(
            "Verluste durch Pleite oder Betrug ("
            + ", ".join(data["lost"]["assets"])
            + ") sind nicht Teil dieser Kette: verlorene Bestände tragen zu keinem späteren "
            "Betrag bei."
        )
    notes.append("Auswertung aus den erfassten Daten, keine Rechts- oder Steuerberatung.")
    for n in notes:
        doc.multi_cell(0, 4.2, f"- {n}")
        doc.ln(0.8)
    flows = data.get("diagram") or {}
    flow = flows.get("simple")
    if flow:
        from app.modules.tax.origin_flow_pdf import draw

        draw(
            doc,
            flow,
            lambda v: f"{_cell(Decimal(str(round(v, 2))), 'eur')} EUR",
            float(Decimal(data["target"]["amount"])),
        )
    if flows.get("full"):
        draw(
            doc,
            flows["full"],
            lambda v: f"{_cell(Decimal(str(round(v, 2))), 'eur')} EUR",
            float(Decimal(data["target"]["amount"])),
            full=True,
        )
    return bytes(doc.output())


COLUMNS = [
    ("Datum", "dt", 28),
    ("Vorgang", "text", 40),
    ("Ort", "text", 50),
    ("Ziel", "text", 40),
    ("Coin", "text", 16),
    ("Menge", "qty", 30),
    ("Wert des Schritts €", "eur", 26),
    ("davon für diesen Betrag €", "eur", 30),
    ("Art", "text", 20),
    ("Transaktions-ID", "text", 60),
    ("Datenart", "text", 14),
]


def rows(data: dict) -> list[list]:
    return [
        [
            _when(s["at"]),
            KIND.get(s["kind"], s["kind"]),
            s["place"] or "außerhalb",
            s.get("to") or "",
            s["asset"],
            _d(s["qty"]),
            _d(s["value"]),
            _d(s["share"]),
            s.get("detail") or "",
            s.get("tx") or "",
            CHANNEL.get(s.get("channel") or "", ""),
        ]
        for s in sorted(data["steps"], key=lambda s: s["at"])
    ]


def xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Zusammenfassung"
    for label, value in head(data):
        ws.append([label, value])
    ws.append([])
    ws.append(["Zusammenfassung", summary(data)])
    ws.append([])
    ws.append(["Herkunft", "Betrag €", "Anteil %"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for m in data["mix"]:
        ws.append([KIND.get(m["kind"], m["kind"]), _d(m["value"]), _d(m["share"])])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30
    if data.get("coins"):
        c = wb.create_sheet("Bestand je Coin")
        kinds = [k for k in KIND if any(k in x["mix"] for x in data["coins"])]
        c.append(["Coin", "Menge", "Wert €", "Kosten €", *[f"{KIND[k]} €" for k in kinds]])
        for cell in c[1]:
            cell.font = Font(bold=True)
        for x in data["coins"]:
            c.append(
                [
                    x["asset"],
                    _d(x["qty"]),
                    _d(x["value"]),
                    _d(x["cost"]),
                    *[_d(x["mix"].get(k)) for k in kinds],
                ]
            )
        c.freeze_panes = "A2"
    if data.get("deposits"):
        d = wb.create_sheet("Einzahlungen")
        d.append(["Datum", "Ort", "Art", "Betrag €", "davon hier €", "Transaktions-ID", "Datenart"])
        for cell in d[1]:
            cell.font = Font(bold=True)
        for x in sorted(data["deposits"], key=lambda x: x["at"]):
            d.append(
                [
                    _when(x["at"]),
                    x["place"] or "außerhalb",
                    PAID[(x["kind"], x.get("detail") == "direct")],
                    _d(x["value"]),
                    _d(x["share"]),
                    x.get("tx") or "",
                    CHANNEL.get(x.get("channel") or "", ""),
                ]
            )
        for (cell,) in d.iter_rows(min_row=2, min_col=1, max_col=1):
            cell.number_format = "dd.mm.yyyy hh:mm"
    ov = data.get("overview") or {}
    if ov.get("sources"):
        q = wb.create_sheet("Datenquellen")
        q.append(["Name", "Typ", "Adresse", "Vorgänge", *[CHANNEL[k] for k in CHANNEL]])
        for cell in q[1]:
            cell.font = Font(bold=True)
        for x in ov["sources"]:
            q.append(
                [
                    x["place"] or "außerhalb",
                    SOURCE_KIND.get(x["kind"], x["kind"]),
                    x.get("address") or "",
                    x["count"],
                    *[x["channels"].get(k, 0) for k in CHANNEL],
                ]
            )
    w = wb.create_sheet("Vorgänge")
    w.append([c[0] for c in COLUMNS])
    for c in w[1]:
        c.font = Font(bold=True)
    for r in rows(data):
        w.append(r)
    for (cell,) in w.iter_rows(min_row=2, min_col=1, max_col=1):
        cell.number_format = "dd.mm.yyyy hh:mm"
    w.freeze_panes = "A2"
    for sheet in wb.worksheets:
        as_text(sheet)  # outside text never becomes a formula (review 2026-10-24)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
