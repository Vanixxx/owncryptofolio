"""Exports as tables (M7): the transactions exactly as the list shows them (same filters,
same order) and the holdings, as Excel or CSV. Derived, nothing stored.

CSV follows the UI language: German with `;` and a decimal comma, so Excel opens it
right; English with `,` and a point. Excel files carry real numbers."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.sheets import as_text, csv_text

LABELS = {
    "de": {
        "tx": [
            "Datum",
            "Art",
            "Quelle",
            "Ziel",
            "Zugang",
            "Coin Zugang",
            "Abgang",
            "Coin Abgang",
            "Gebühr",
            "Coin Gebühr",
            "Netzgebühr",
            "Wert ({cur})",
            "Hinweis",
            "Tx-ID",
            "Kommentar",
        ],
        "holdings": [
            "Coin",
            "Name",
            "Menge",
            "DCA ({cur})",
            "Kurs ({cur})",
            "Einstand ({cur})",
            "Wert ({cur})",
            "Unrealisiert ({cur})",
            "Unrealisiert %",
            "Realisiert ({cur})",
            "Erträge ({cur})",
            "Im Besitz seit",
            "Orte",
        ],
        "kinds": {
            "trade": "Handel",
            "send": "Senden",
            "receive": "Empfangen",
            "fiat_deposit": "Fiat-Einzahlung",
            "fiat_withdrawal": "Fiat-Auszahlung",
            "income": "Ertrag",
            "internal_transfer": "Umbuchung",
            "migration": "Umstellung",
            "fee": "Gebühr",
            "realized_pnl": "Realisiertes Ergebnis",
            "funding": "Funding",
            "futures_position": "Futures-Position",
            "other": "Sonstiges",
            "transfer": "Transfer",
        },
        "notes": {
            "open": "kein Gegenstück",
            "ambiguous": "mehrdeutig",
            "dust": "Staub?",
            "spam": "Spam",
            "external": "an Fremde",
            "own_untracked": "eigene Wallet, nicht erfasst",
            "lost": "verloren",
            "bought": "außerhalb gekauft",
            "income": "Ertrag",
            "outside": "eigene Wallet, nicht erfasst",
        },
        "sheet_tx": "Transaktionen",
        "sheet_holdings": "Bestände",
    },
    "en": {
        "tx": [
            "Date",
            "Kind",
            "Source",
            "Target",
            "In",
            "In coin",
            "Out",
            "Out coin",
            "Fee",
            "Fee coin",
            "Network fee",
            "Value ({cur})",
            "Note",
            "Tx id",
            "Comment",
        ],
        "holdings": [
            "Coin",
            "Name",
            "Quantity",
            "DCA ({cur})",
            "Price ({cur})",
            "Cost ({cur})",
            "Value ({cur})",
            "Unrealized ({cur})",
            "Unrealized %",
            "Realized ({cur})",
            "Income ({cur})",
            "Held since",
            "Places",
        ],
        "kinds": {
            "trade": "Trade",
            "send": "Send",
            "receive": "Receive",
            "fiat_deposit": "Fiat deposit",
            "fiat_withdrawal": "Fiat withdrawal",
            "income": "Income",
            "internal_transfer": "Internal transfer",
            "migration": "Migration",
            "fee": "Fee",
            "realized_pnl": "Realized result",
            "funding": "Funding",
            "futures_position": "Futures position",
            "other": "Other",
            "transfer": "Transfer",
        },
        "notes": {
            "open": "no counterpart",
            "ambiguous": "ambiguous",
            "dust": "dust?",
            "spam": "spam",
            "external": "to others",
            "own_untracked": "own wallet, not tracked",
            "lost": "lost",
            "bought": "bought outside",
            "income": "income",
            "outside": "own wallet, not tracked",
        },
        "sheet_tx": "Transactions",
        "sheet_holdings": "Holdings",
    },
}


def _d(v) -> Decimal | None:
    return None if v in (None, "") else Decimal(str(v))


def transaction_rows(items: list[dict], values: dict, name, lang: str, tz: str) -> list[list]:
    """One table row per list line: a transfer is one row (source → target), a futures
    position one row with its result."""
    L = LABELS[lang]
    zone = ZoneInfo(tz)
    when = lambda iso: datetime.fromisoformat(iso).astimezone(zone).replace(tzinfo=None)  # noqa: E731
    out = []
    for i in items:
        value = values.get(id(i))
        if i["type"] == "transfer":
            s, r = i["send"], i["receive"]
            net = i["network_fee"] or [None, None]
            out.append(
                [
                    when(s["timestamp"]),
                    L["kinds"]["transfer"],
                    name(s["source_id"]),
                    name(r["source_id"]),
                    _d(r["incoming"][1]),
                    r["incoming"][0],
                    _d(s["outgoing"][1]),
                    s["outgoing"][0],
                    _d(s["fee"][1]),
                    s["fee"][0],
                    _d(net[1]),
                    value,
                    "",
                    s["external_id"],
                    s["comment"],
                ]
            )
        elif i["type"] == "position":
            res = _d(i["result"]) or Decimal(0)
            note = f"{i['instrument']} {i['side'] or ''}".strip()
            out.append(
                [
                    when(i["closed"]),
                    L["kinds"]["futures_position"],
                    name(i["source_id"]),
                    "",
                    res if res > 0 else None,
                    i["coin"] if res > 0 else None,
                    -res if res < 0 else None,
                    i["coin"] if res < 0 else None,
                    _d(i["fees"]),
                    i["coin"],
                    None,
                    value,
                    note,
                    None,
                    f"pnl {i['pnl']} · funding {i['funding']}",
                ]
            )
        else:
            row = i["row"]
            notes = [L["notes"][k] for k in ("open", "ambiguous", "dust") if i.get(k)]
            if i.get("spam"):
                notes.append(L["notes"]["spam"])
            if i.get("open_mark"):
                notes = [L["notes"][i["open_mark"]]]
            kind = L["kinds"].get(row["kind"], row["kind"])
            if row["market"] == "futures":
                kind = f"{kind} · {row['instrument'] or 'Futures'}"
            out.append(
                [
                    when(row["timestamp"]),
                    kind,
                    name(row["source_id"]),
                    "",
                    _d(row["incoming"][1]),
                    row["incoming"][0],
                    _d(row["outgoing"][1]),
                    row["outgoing"][0],
                    _d(row["fee"][1]),
                    row["fee"][0],
                    None,
                    value,
                    ", ".join(notes),
                    row["external_id"],
                    row["comment"],
                ]
            )
    return out


def _sig(v: Decimal | None, digits: int = 8) -> Decimal | None:
    """A price for reading: eight significant digits (0.1089933681… → 0.10899337)."""
    if v is None or v == 0:
        return v
    return round(v, max(0, digits - v.adjusted() - 1)).normalize()


def holding_rows(items: list[dict], name, lang: str) -> list[list]:
    L = LABELS[lang]
    out = []
    for h in items:
        if h["fiat"] and not _d(h["qty"]):
            continue
        places = ", ".join(
            f"{L['notes']['outside'] if p['outside'] else name(p['source_id'])}: "
            + (p["qty"].replace(".", ",") if lang == "de" else p["qty"])
            for p in h["places"]
        )
        out.append(
            [
                h["asset"],
                h["name"],
                _d(h["qty"]),
                _sig(_d(h["avg_price"])),
                _sig(_d(h["price"])),
                _d(h["cost"]),
                _d(h["value"]),
                _d(h["unrealized"]),
                _d(h["unrealized_pct"]),
                _d(h["realized"]),
                _d(h["income"]),
                (h["since"] or "")[:10] or None,
                places,
            ]
        )
    return out


def to_csv(header: list[str], rows: list[list], lang: str) -> bytes:
    sep, comma = (";", True) if lang == "de" else (",", False)

    def cell(v):
        if v is None:
            return ""
        if isinstance(v, Decimal):
            text = format(v.normalize(), "f")
            return text.replace(".", ",") if comma else text
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return csv_text(str(v))

    buf = io.StringIO()
    w = csv.writer(buf, delimiter=sep, lineterminator="\r\n")
    w.writerow(header)
    for r in rows:
        w.writerow([cell(v) for v in r])
    # BOM: Excel otherwise reads UTF-8 as ANSI and breaks umlauts
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(sheet: str, header: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for col, title in enumerate(header, 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = max(10, min(40, len(title) + 4))
        for cell in ws[letter][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
    ws.auto_filter.ref = ws.dimensions
    as_text(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
